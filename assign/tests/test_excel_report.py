#!/usr/bin/env python3
"""
Unit tests for excel_report.py — style_header_row's header/autofilter/
freeze-pane setup and autofit_columns' width clamping.

Run with: python3 -m unittest discover -s assign/tests -v
"""

import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from excel_report import autofit_columns, style_header_row


class TestStyleHeaderRow(unittest.TestCase):
    def test_writes_header_row(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        style_header_row(ws, ["A", "B", "C"])
        self.assertEqual([c.value for c in ws[1]], ["A", "B", "C"])

    def test_header_cells_are_bold_and_filled(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        style_header_row(ws, ["A", "B"])
        for cell in ws[1]:
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.fill.fgColor.rgb, "00BDD7EE")

    def test_enables_autofilter_and_freeze_panes(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        style_header_row(ws, ["A", "B"])
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertIsNotNone(ws.auto_filter.ref)

    def test_custom_font_and_fill_override_defaults(self):
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        custom_font = Font(bold=True, size=20)
        custom_fill = PatternFill("solid", fgColor="FF0000")
        style_header_row(ws, ["A"], font=custom_font, fill=custom_fill)
        self.assertEqual(ws["A1"].font.size, 20)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "00FF0000")


class TestAutofitColumns(unittest.TestCase):
    def test_clamps_to_min_width(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["x"])
        autofit_columns(ws, min_width=20, max_width=55)
        self.assertEqual(ws.column_dimensions["A"].width, 20)

    def test_clamps_to_max_width(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["x" * 200])
        autofit_columns(ws, min_width=12, max_width=55)
        self.assertEqual(ws.column_dimensions["A"].width, 55)

    def test_fits_content_between_bounds(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["12345678901234567890"])   # 20 chars
        autofit_columns(ws, min_width=12, max_width=55)
        self.assertEqual(ws.column_dimensions["A"].width, 22)   # len + 2

    def test_uses_header_row_when_longer_than_data(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A Longer Header Name"])
        ws.append(["short"])
        autofit_columns(ws, min_width=12, max_width=55)
        self.assertEqual(ws.column_dimensions["A"].width, len("A Longer Header Name") + 2)


if __name__ == "__main__":
    unittest.main()
