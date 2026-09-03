#!/usr/bin/env python3
"""
Unit tests for task_analytics.py — timestamp normalization, the merged-
record builder, every milestone function, list-fetch pagination/dedup,
detail-fetch branch selection (the valuation_request_id shortcut and the
registrar backfill search), the Excel/Telegram output builders, and the
conversation handlers.

Run with: python3 -m unittest discover -s assign/tests -v
"""

import asyncio
import os
import sys
import unittest
from datetime import datetime, timedelta
from unittest.mock import AsyncMock, MagicMock, patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import task_analytics as ta
from token_rotator import _AllTokensExhausted, _TokenRotator


def _run(coro):
    return asyncio.run(coro)


def _make_query_update(data):
    update = MagicMock()
    query = update.callback_query
    query.data = data
    query.answer = AsyncMock()
    query.edit_message_text = AsyncMock()
    query.message.reply_text = AsyncMock()
    return update


def _make_message_update(text):
    update = MagicMock()
    update.message.text = text
    update.message.reply_text = AsyncMock()
    return update


# ──────────────────────────────────────────────────────────
# _ta_parse_ts
# ──────────────────────────────────────────────────────────
class TestParseTs(unittest.TestCase):
    def test_explicit_utc_z_converts_to_naive_local(self):
        # 14:09:57Z (UTC) -> 17:09:57 local (EAT, UTC+3)
        result = ta._ta_parse_ts("2025-12-09T14:09:57.239167Z")
        self.assertEqual(result, datetime(2025, 12, 9, 17, 9, 57, 239167))

    def test_explicit_offset_converts_to_naive_local(self):
        result = ta._ta_parse_ts("2025-12-09T17:25:43.987022+03:00")
        self.assertEqual(result, datetime(2025, 12, 9, 17, 25, 43, 987022))

    def test_naive_iso_string_passes_through_unchanged(self):
        result = ta._ta_parse_ts("2025-12-09T17:25:43.193553")
        self.assertEqual(result, datetime(2025, 12, 9, 17, 25, 43, 193553))

    def test_naive_space_separated_actors_style_parses(self):
        result = ta._ta_parse_ts("2025-12-09 17:31:42.682292")
        self.assertEqual(result, datetime(2025, 12, 9, 17, 31, 42, 682292))

    def test_literal_none_string_returns_none(self):
        self.assertIsNone(ta._ta_parse_ts("None"))

    def test_actual_none_returns_none(self):
        self.assertIsNone(ta._ta_parse_ts(None))

    def test_empty_string_returns_none(self):
        self.assertIsNone(ta._ta_parse_ts(""))

    def test_garbage_string_returns_none(self):
        self.assertIsNone(ta._ta_parse_ts("not-a-date"))


# ──────────────────────────────────────────────────────────
# _ta_build_merged_record
# ──────────────────────────────────────────────────────────
class TestBuildMergedRecord(unittest.TestCase):
    def test_registrar_only(self):
        registrar = {
            "officers": [{"names": "X", "role": "COUNTY_REGISTRAR", "assigned_date": "2025-12-09T14:09:57Z"}],
            "county": "BUNGOMA", "registry": "BUNGOMA", "application_status": "ONGOING",
            "valuation_request_id": None,
        }
        merged = ta._ta_build_merged_record("REF1", registrar, None)
        self.assertEqual(merged["county"], "BUNGOMA")
        self.assertFalse(merged["has_valuation_record"])
        self.assertTrue(merged["has_registrar_record"])
        self.assertEqual(merged["date_created_registrar"], "2025-12-09T14:09:57Z")
        self.assertIsNone(merged["date_created_valuation"])

    def test_valuation_only(self):
        valuation = {"date_created": "2025-12-09T17:25:43", "county": "NAIROBI",
                     "application_status": "ONGOING", "remarks": [], "actors": []}
        merged = ta._ta_build_merged_record("REF1", None, valuation)
        self.assertFalse(merged["has_registrar_record"])
        self.assertTrue(merged["has_valuation_record"])
        self.assertIsNone(merged["date_created_registrar"])
        self.assertEqual(merged["date_created_valuation"], "2025-12-09T17:25:43")

    def test_both_present_prefers_valuation_county(self):
        registrar = {"officers": [], "county": "BUNGOMA", "application_status": "ONGOING"}
        valuation = {"date_created": "2025-12-09T17:25:43", "county": "NAIROBI",
                     "application_status": "COMPLETED", "remarks": [], "actors": []}
        merged = ta._ta_build_merged_record("REF1", registrar, valuation)
        self.assertEqual(merged["county"], "NAIROBI")
        self.assertEqual(merged["application_status"], "COMPLETED")

    def test_no_officers_registrar_date_is_none(self):
        merged = ta._ta_build_merged_record("REF1", {"officers": []}, None)
        self.assertIsNone(merged["date_created_registrar"])


# ──────────────────────────────────────────────────────────
# Milestone functions — built against the real completed-task sample
# (CNTYINV/5MJ9KT6J1P) from the conversation that shaped this module.
# ──────────────────────────────────────────────────────────
def _completed_merged():
    registrar_detail = {
        "officers": [{"names": "ABDUL JUNIOR NELFRANK", "role": "COUNTY_REGISTRAR",
                      "assigned_date": "2025-12-09T14:09:57.239167Z"}],
        "county": "BUNGOMA", "registry": "BUNGOMA", "application_status": "COMPLETED",
        "valuation_request_id": "57997a5d-d378-4637-aa26-f319885b0d18",
    }
    valuation_detail = {
        "date_created": "2025-12-09T17:25:43.193553",
        "application_status": "COMPLETED",
        "county": "BUNGOMA", "registry": "BUNGOMA",
        "remarks": [
            {"status": "ASSESSOR APPROVAL REMARKS", "date_created": "2025-12-09T17:25:43.354618"},
            {"status": "DLV_FORWARDING_REMARKS", "date_created": "2025-12-09T17:31:42.712145"},
            {"status": "VO_FORWARD_REMARKS", "date_created": "2025-12-09T17:35:34.394535"},
            {"status": "DLV_APPROVAL_REMARKS", "date_created": "2025-12-09T17:36:20.214399"},
        ],
        "actors": [
            {"role": "VALUATION OFFICER",
             "user_details": {"id": "b5a4bf31", "names": "JOSPHAT MWENYEWE MWANA"},
             "date_assigned": "2025-12-09 17:31:42.682292"},
            {"role": "VALUER_IN_CHARGE",
             "user_details": {"id": "b5a4bf31", "names": "JOSPHAT MWENYEWE MWANA"},
             "date_assigned": "None"},
        ],
    }
    return ta._ta_build_merged_record("CNTYINV/5MJ9KT6J1P", registrar_detail, valuation_detail)


class TestCreatedAt(unittest.TestCase):
    def test_registrar_assigned_date_used_when_present(self):
        merged = _completed_merged()
        self.assertEqual(ta._ta_created_at(merged), datetime(2025, 12, 9, 17, 9, 57, 239167))

    def test_falls_back_to_valuation_date_created_when_no_registrar_record(self):
        merged = ta._ta_build_merged_record("R", None, {"date_created": "2025-12-09T17:25:43"})
        self.assertEqual(ta._ta_created_at(merged), datetime(2025, 12, 9, 17, 25, 43))

    def test_none_when_neither_present(self):
        merged = ta._ta_build_merged_record("R", None, None)
        self.assertIsNone(ta._ta_created_at(merged))


class TestTimeAtRegistrar(unittest.TestCase):
    def test_valuation_request_id_populated_gives_completed_duration(self):
        merged = _completed_merged()
        duration, ongoing = ta._ta_time_at_registrar(merged, datetime(2026, 1, 1))
        self.assertFalse(ongoing)
        self.assertAlmostEqual(duration.total_seconds(), 945.954386, places=2)

    def test_null_valuation_request_id_gives_ongoing_elapsed(self):
        merged = ta._ta_build_merged_record(
            "R", {"officers": [{"assigned_date": "2025-01-01T00:00:00Z"}], "valuation_request_id": None}, None)
        now = datetime(2025, 1, 4, 3, 0, 0)
        duration, ongoing = ta._ta_time_at_registrar(merged, now)
        self.assertTrue(ongoing)
        self.assertEqual(duration, timedelta(days=3))

    def test_unresolvable_created_at_returns_none(self):
        merged = ta._ta_build_merged_record("R", None, None)
        duration, ongoing = ta._ta_time_at_registrar(merged, datetime(2026, 1, 1))
        self.assertIsNone(duration)
        self.assertFalse(ongoing)

    def test_populated_id_but_unresolvable_entered_returns_none_not_ongoing(self):
        merged = ta._ta_build_merged_record(
            "R", {"officers": [{"assigned_date": "2025-01-01T00:00:00Z"}],
                  "valuation_request_id": "abc"}, None)
        duration, ongoing = ta._ta_time_at_registrar(merged, datetime(2026, 1, 1))
        self.assertIsNone(duration)
        self.assertFalse(ongoing)


class TestRemarksByStatus(unittest.TestCase):
    def test_filters_and_sorts_ascending(self):
        merged = _completed_merged()
        times = ta._ta_remarks_by_status(merged, "DLV_FORWARDING_REMARKS")
        self.assertEqual(times, [datetime(2025, 12, 9, 17, 31, 42, 712145)])

    def test_missing_remarks_returns_empty(self):
        merged = ta._ta_build_merged_record("R", None, None)
        self.assertEqual(ta._ta_remarks_by_status(merged, "X"), [])

    def test_no_matching_status_returns_empty(self):
        merged = _completed_merged()
        self.assertEqual(ta._ta_remarks_by_status(merged, "NOPE"), [])


class TestFirstAndCurrentValuerAssignedAt(unittest.TestCase):
    def test_first_valuer_from_forwarding_remark(self):
        merged = _completed_merged()
        self.assertEqual(ta._ta_first_valuer_assigned_at(merged),
                          datetime(2025, 12, 9, 17, 31, 42, 712145))

    def test_first_valuer_falls_back_to_actors_when_no_remarks(self):
        valuation = {
            "date_created": "2025-01-01T00:00:00", "remarks": [],
            "actors": [{"role": "VALUATION OFFICER", "date_assigned": "2025-01-02T00:00:00"}],
        }
        merged = ta._ta_build_merged_record("R", None, valuation)
        self.assertEqual(ta._ta_first_valuer_assigned_at(merged), datetime(2025, 1, 2))

    def test_current_valuer_prefers_actors_over_remarks_when_both_resolve(self):
        merged = _completed_merged()
        # actors[].date_assigned (17:31:42.682292) differs slightly from the
        # DLV_FORWARDING_REMARKS timestamp (17:31:42.712145) — actors wins.
        self.assertEqual(ta._ta_current_valuer_assigned_at(merged),
                          datetime(2025, 12, 9, 17, 31, 42, 682292))

    def test_current_valuer_falls_back_to_remarks_when_actors_unresolvable(self):
        valuation = {
            "date_created": "2025-01-01T00:00:00",
            "remarks": [{"status": "DLV_REASSIGN_REMARKS", "date_created": "2025-01-03T00:00:00"}],
            "actors": [{"role": "VALUATION OFFICER", "date_assigned": "None"}],
        }
        merged = ta._ta_build_merged_record("R", None, valuation)
        self.assertEqual(ta._ta_current_valuer_assigned_at(merged), datetime(2025, 1, 3))

    def test_none_when_neither_resolves(self):
        merged = ta._ta_build_merged_record("R", None, None)
        self.assertIsNone(ta._ta_first_valuer_assigned_at(merged))
        self.assertIsNone(ta._ta_current_valuer_assigned_at(merged))


class TestHandoffCount(unittest.TestCase):
    def test_counts_only_reassign_remarks(self):
        merged = _completed_merged()  # no DLV_REASSIGN_REMARKS in this sample
        self.assertEqual(ta._ta_handoff_count(merged), 0)

    def test_counts_multiple_reassignments(self):
        valuation = {"remarks": [
            {"status": "DLV_REASSIGN_REMARKS", "date_created": "2025-01-01T00:00:00"},
            {"status": "DLV_REASSIGN_REMARKS", "date_created": "2025-01-02T00:00:00"},
            {"status": "DLV_FORWARDING_REMARKS", "date_created": "2025-01-01T00:00:00"},
        ]}
        merged = ta._ta_build_merged_record("R", None, valuation)
        self.assertEqual(ta._ta_handoff_count(merged), 2)


class TestFinalizedAt(unittest.TestCase):
    def test_last_remark_regardless_of_status_when_completed(self):
        merged = _completed_merged()
        self.assertEqual(ta._ta_finalized_at(merged), datetime(2025, 12, 9, 17, 36, 20, 214399))

    def test_none_when_not_completed(self):
        merged = ta._ta_build_merged_record(
            "R", None, {"application_status": "ONGOING",
                        "remarks": [{"status": "X", "date_created": "2025-01-01T00:00:00"}]})
        self.assertIsNone(ta._ta_finalized_at(merged))

    def test_none_when_completed_but_no_remarks(self):
        merged = ta._ta_build_merged_record("R", None, {"application_status": "COMPLETED", "remarks": []})
        self.assertIsNone(ta._ta_finalized_at(merged))


class TestTotalAgeAndTimeWithCurrentValuer(unittest.TestCase):
    def test_total_age_completed(self):
        merged = _completed_merged()
        duration, ongoing = ta._ta_total_age(merged, datetime(2026, 1, 1))
        self.assertFalse(ongoing)
        self.assertAlmostEqual(duration.total_seconds(), 1582.975232, places=2)

    def test_total_age_ongoing(self):
        merged = ta._ta_build_merged_record(
            "R", {"officers": [{"assigned_date": "2025-01-01T00:00:00Z"}]}, None)
        now = datetime(2025, 1, 3, 3, 0, 0)
        duration, ongoing = ta._ta_total_age(merged, now)
        self.assertTrue(ongoing)
        self.assertEqual(duration, timedelta(days=2))

    def test_total_age_unresolvable_created_returns_none(self):
        merged = ta._ta_build_merged_record("R", None, None)
        duration, ongoing = ta._ta_total_age(merged, datetime(2026, 1, 1))
        self.assertIsNone(duration)
        self.assertFalse(ongoing)

    def test_time_with_current_valuer_completed(self):
        merged = _completed_merged()
        duration, ongoing = ta._ta_time_with_current_valuer(merged, datetime(2026, 1, 1))
        self.assertFalse(ongoing)
        self.assertAlmostEqual(duration.total_seconds(), 277.532107, places=2)

    def test_time_with_current_valuer_none_when_never_reached_dlv(self):
        merged = ta._ta_build_merged_record("R", None, None)
        duration, ongoing = ta._ta_time_with_current_valuer(merged, datetime(2026, 1, 1))
        self.assertIsNone(duration)
        self.assertFalse(ongoing)


class TestStageAndValuerHelpers(unittest.TestCase):
    def test_stage_label_completed(self):
        self.assertEqual(ta._ta_current_stage_label(_completed_merged()), "Completed")

    def test_stage_label_at_dlv_with_valuer(self):
        merged = ta._ta_build_merged_record(
            "R", None, {"application_status": "ONGOING",
                        "actors": [{"role": "VALUATION OFFICER", "user_details": {"names": "Jane"}}]})
        self.assertEqual(ta._ta_current_stage_label(merged), "DLV — Jane")

    def test_stage_label_at_dlv_unassigned(self):
        merged = ta._ta_build_merged_record("R", None, {"application_status": "ONGOING", "actors": []})
        self.assertEqual(ta._ta_current_stage_label(merged), "DLV — unassigned")

    def test_stage_label_registrar_only(self):
        merged = ta._ta_build_merged_record("R", {"officers": []}, None)
        self.assertEqual(ta._ta_current_stage_label(merged), "Registrar")

    def test_current_valuer_name_and_id(self):
        merged = _completed_merged()
        self.assertEqual(ta._ta_current_valuer_name(merged), "JOSPHAT MWENYEWE MWANA")
        self.assertEqual(ta._ta_current_valuer_id(merged), "b5a4bf31")


class TestComputeMilestones(unittest.TestCase):
    def test_smoke_all_keys_present(self):
        merged = _completed_merged()
        milestones = ta._ta_compute_milestones(merged, datetime(2026, 1, 1))
        for key in ("created_at", "entered_valuation_at", "time_at_registrar", "registrar_ongoing",
                    "first_valuer_assigned_at", "current_valuer_assigned_at", "time_at_valuer",
                    "valuer_stage_ongoing", "time_with_current_valuer", "current_valuer_ongoing",
                    "handoff_count", "finalized_at", "total_age", "age_ongoing",
                    "current_valuer_name", "current_valuer_id", "stage_label"):
            self.assertIn(key, milestones)
        self.assertEqual(milestones["stage_label"], "Completed")
        self.assertEqual(milestones["handoff_count"], 0)


# ──────────────────────────────────────────────────────────
# List fetch — pagination, cutoff early-stop, dedup
# ──────────────────────────────────────────────────────────
class TestFetchRegistrarList(unittest.TestCase):
    def test_stops_on_cutoff_and_dedupes_across_filters(self):
        http_sess = MagicMock()
        page1 = {"results": [
            {"reference_number": "R1", "date_created": "2026-01-05"},
            {"reference_number": "R2", "date_created": "2025-01-01"},  # below cutoff
        ], "next": None}
        http_sess.get.return_value = MagicMock(
            raise_for_status=lambda: None, json=lambda: page1)
        with patch.object(ta, "_ta_fetch_registrar_page", side_effect=lambda *a, **k: page1):
            results = ta._ta_fetch_registrar_list(http_sess, MagicMock(), "2026-01-01")
        refs = {r["reference_number"] for r in results}
        self.assertEqual(refs, {"R1"})

    def test_pagination_stops_when_no_next(self):
        with patch.object(ta, "_ta_fetch_registrar_page",
                           return_value={"results": [{"reference_number": "R1", "date_created": "2026-01-05"}],
                                         "next": None}) as mock_page:
            ta._ta_fetch_registrar_list(MagicMock(), MagicMock(), "2026-01-01")
        # Called once per filter (2 filters), page=1 each — no second page fetched.
        self.assertEqual(mock_page.call_count, len(ta._TA_REGISTRAR_FILTERS))

    def test_empty_results_breaks_immediately(self):
        with patch.object(ta, "_ta_fetch_registrar_page", return_value={"results": [], "next": None}):
            results = ta._ta_fetch_registrar_list(MagicMock(), MagicMock(), "2026-01-01")
        self.assertEqual(results, [])

    def test_page_error_breaks_that_filters_loop(self):
        with patch.object(ta, "_ta_fetch_registrar_page", side_effect=Exception("boom")):
            results = ta._ta_fetch_registrar_list(MagicMock(), MagicMock(), "2026-01-01")
        self.assertEqual(results, [])


class TestFetchValuationList(unittest.TestCase):
    def test_dedupes_by_reference_number(self):
        page = {"results": [{"reference_number": "R1", "date_created": "2026-01-05"}], "next": None}
        with patch.object(ta, "_ta_fetch_valuation_page", return_value=page):
            results = ta._ta_fetch_valuation_list(MagicMock(), MagicMock(), "2026-01-01")
        # Same ref appears in both "Ongoing" and "Completed" filter passes — deduped to one.
        self.assertEqual(len(results), 1)


# ──────────────────────────────────────────────────────────
# _ta_merge_populations
# ──────────────────────────────────────────────────────────
class TestMergePopulations(unittest.TestCase):
    def test_merges_by_ref_no_filter(self):
        registrar = [{"reference_number": "R1", "county": "NAIROBI"}]
        valuation = [{"reference_number": "R1", "county": "NAIROBI"}, {"reference_number": "R2", "county": "KISUMU"}]
        merged = ta._ta_merge_populations(registrar, valuation, "")
        self.assertEqual(set(merged.keys()), {"R1", "R2"})
        self.assertIn("registrar_item", merged["R1"])
        self.assertIn("valuation_item", merged["R1"])
        self.assertNotIn("registrar_item", merged["R2"])

    def test_county_filter_prefers_valuation_side_county(self):
        registrar = [{"reference_number": "R1", "county": "BUNGOMA"}]
        valuation = [{"reference_number": "R1", "county": "NAIROBI"}]
        merged = ta._ta_merge_populations(registrar, valuation, "nairobi")
        self.assertIn("R1", merged)

    def test_county_filter_excludes_non_matching(self):
        registrar = [{"reference_number": "R1", "county": "KISUMU"}]
        merged = ta._ta_merge_populations(registrar, [], "nairobi")
        self.assertEqual(merged, {})


# ──────────────────────────────────────────────────────────
# Detail-fetch wrappers + _ta_process_ref branch selection
# ──────────────────────────────────────────────────────────
class TestFetchDetailWrappers(unittest.TestCase):
    def test_registrar_detail_unwraps_details_key(self):
        with patch.object(ta, "fetch_with_rotation", return_value={"details": {"node": "X"}}):
            result = ta._ta_fetch_registrar_detail(MagicMock(), MagicMock(), "id-1")
        self.assertEqual(result, {"node": "X"})

    def test_valuation_detail_returns_raw_json(self):
        with patch.object(ta, "fetch_with_rotation", return_value={"node": "Y"}):
            result = ta._ta_fetch_valuation_detail(MagicMock(), MagicMock(), "id-1")
        self.assertEqual(result, {"node": "Y"})


class TestProcessRef(unittest.TestCase):
    def test_valuation_request_id_shortcut_skips_independent_search(self):
        """When the registrar detail already links to a valuation record,
        fetch it directly — never call the county DLV search primitives."""
        wrapped_registrar_detail = {"details": {"officers": [], "valuation_request_id": "vid-1",
                                                  "application_status": "ONGOING"}}
        with patch.object(ta, "fetch_with_rotation",
                           side_effect=[wrapped_registrar_detail, {"node": "Y", "date_created": "2025-01-01T00:00:00"}]), \
             patch.object(ta, "_lu_search_ref_county") as mock_search:
            row = ta._ta_process_ref(
                MagicMock(), _TokenRotator([("staff2", MagicMock())]), _TokenRotator([("staff_valuer", MagicMock())]),
                "REF1", {"id": "reg-id"}, None, datetime(2026, 1, 1),
            )
        mock_search.assert_not_called()
        self.assertTrue(row["merged"]["has_valuation_record"])

    def test_valuation_list_item_id_used_when_no_valuation_request_id(self):
        wrapped_registrar_detail = {"details": {"officers": [], "valuation_request_id": None,
                                                  "application_status": "ONGOING"}}
        with patch.object(ta, "fetch_with_rotation",
                           side_effect=[wrapped_registrar_detail, {"node": "Y", "date_created": "2025-01-01T00:00:00"}]):
            row = ta._ta_process_ref(
                MagicMock(), _TokenRotator([("staff2", MagicMock())]), _TokenRotator([("staff_valuer", MagicMock())]),
                "REF1", {"id": "reg-id"}, {"id": "val-id"}, datetime(2026, 1, 1),
            )
        self.assertTrue(row["merged"]["has_valuation_record"])

    def test_backfill_search_runs_when_only_found_via_valuation_list(self):
        """No registrar_item at all — recover the registrar record with an
        unbounded single-ref search rather than losing time_at_registrar."""
        registrar_tokens = MagicMock()
        with patch.object(ta, "fetch_with_rotation", return_value={"node": "Y", "date_created": "2025-01-01T00:00:00"}), \
             patch.object(ta, "_lu_search_ref_county", return_value={"id": "backfill-id"}) as mock_search, \
             patch.object(ta, "_lu_fetch_detail_county", return_value={"officers": [], "application_status": "ONGOING"}) as mock_detail:
            row = ta._ta_process_ref(
                MagicMock(), _TokenRotator([("staff2", registrar_tokens)]), _TokenRotator([("staff_valuer", MagicMock())]),
                "REF1", None, {"id": "val-id"}, datetime(2026, 1, 1),
            )
        mock_search.assert_called_once_with(registrar_tokens, "REF1")
        mock_detail.assert_called_once_with(registrar_tokens, "backfill-id")
        self.assertTrue(row["merged"]["has_registrar_record"])

    def test_all_tokens_exhausted_propagates(self):
        with patch.object(ta, "fetch_with_rotation", side_effect=_AllTokensExhausted("boom")):
            with self.assertRaises(_AllTokensExhausted):
                ta._ta_process_ref(
                    MagicMock(), _TokenRotator([("staff2", MagicMock())]), _TokenRotator([("staff_valuer", MagicMock())]),
                    "REF1", {"id": "reg-id"}, None, datetime(2026, 1, 1),
                )


# ──────────────────────────────────────────────────────────
# Excel builder
# ──────────────────────────────────────────────────────────
class TestBuildExcel(unittest.TestCase):
    def _row(self, ref="REF1", status="COMPLETED"):
        merged = _completed_merged()
        merged["ref"] = ref
        milestones = ta._ta_compute_milestones(merged, datetime(2026, 1, 1))
        return {"ref": ref, "merged": merged, "milestones": milestones,
                "county": "BUNGOMA", "registry": "BUNGOMA", "application_status": status}

    def test_sheet_names_present(self):
        wb = openpyxl.load_workbook(__import__("io").BytesIO(ta._ta_build_excel([self._row()])))
        self.assertEqual(wb.sheetnames, ["Task Detail", "Summary"])

    def test_detail_sheet_has_header_and_one_data_row(self):
        wb = openpyxl.load_workbook(__import__("io").BytesIO(ta._ta_build_excel([self._row()])))
        ws = wb["Task Detail"]
        self.assertEqual(ws.cell(row=1, column=1).value, "Reference Number")
        self.assertEqual(ws.cell(row=2, column=1).value, "REF1")

    def test_summary_sheet_has_overall_totals(self):
        wb = openpyxl.load_workbook(__import__("io").BytesIO(ta._ta_build_excel([self._row()])))
        ws = wb["Summary"]
        self.assertEqual(ws.cell(row=1, column=1).value, "Overall Totals")
        self.assertEqual(ws.cell(row=2, column=1).value, "Total Refs")
        self.assertEqual(ws.cell(row=2, column=2).value, 1)


# ──────────────────────────────────────────────────────────
# Telegram field builders
# ──────────────────────────────────────────────────────────
class TestFieldBuilders(unittest.TestCase):
    def setUp(self):
        merged = _completed_merged()
        self.milestones = ta._ta_compute_milestones(merged, datetime(2026, 1, 1))

    def test_stage_field(self):
        self.assertEqual(ta._ta_stage_field(self.milestones), ("🔄 Stage", "Completed"))

    def test_age_field_shows_completed(self):
        label, value = ta._ta_age_field(self.milestones)
        self.assertIn("completed", value)

    def test_registrar_time_field_dash_when_no_registrar_record(self):
        milestones = ta._ta_compute_milestones(ta._ta_build_merged_record("R", None, None), datetime(2026, 1, 1))
        self.assertEqual(ta._ta_registrar_time_field(milestones), ("📋 Time at Registrar", "—"))

    def test_handoff_field_none_when_zero(self):
        self.assertIsNone(ta._ta_handoff_field(self.milestones))

    def test_handoff_field_present_when_nonzero(self):
        milestones = dict(self.milestones)
        milestones["handoff_count"] = 3
        self.assertEqual(ta._ta_handoff_field(milestones), ("🔁 Reassignments", "3"))

    def test_fields_for_appends_handoff_only_when_present(self):
        fields = ta._ta_fields_for(self.milestones)
        self.assertEqual(len(fields), 4)  # no reassignments in this sample


# ──────────────────────────────────────────────────────────
# Handlers
# ──────────────────────────────────────────────────────────
class TestCmdTaskAnalytics(unittest.TestCase):
    def test_both_creds_present_starts_county_step(self):
        update = _make_message_update("/taskanalytics")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(ta, "allowed", return_value=True), \
             patch.object(ta, "get_valid_tokens", return_value=MagicMock()):
            state = _run(ta.cmd_task_analytics(update, ctx))
        self.assertEqual(state, ta.TA.COUNTY)
        self.assertIn("ta_session", ctx.user_data)

    def test_missing_credential_hard_blocks(self):
        update = _make_message_update("/taskanalytics")
        ctx = MagicMock()
        ctx.user_data = {}

        def _tokens(cred_type):
            return None if cred_type == "staff2" else MagicMock()

        with patch.object(ta, "allowed", return_value=True), \
             patch.object(ta, "get_valid_tokens", side_effect=_tokens):
            state = _run(ta.cmd_task_analytics(update, ctx))
        self.assertEqual(state, ta.ConversationHandler.END)
        text = update.message.reply_text.call_args[0][0]
        self.assertIn("Support Reg", text)


class TestRecvTaCounty(unittest.TestCase):
    def test_nairobi_selection_stores_lowercase_and_moves_to_period(self):
        update = _make_query_update("ft_county:nairobi")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(ta, "allowed", return_value=True):
            state = _run(ta.recv_ta_county(update, ctx))
        self.assertEqual(state, ta.TA.PERIOD)
        self.assertEqual(ctx.user_data["ta_session"].county, "nairobi")

    def test_all_selection_stores_empty_string(self):
        update = _make_query_update("ft_county:all")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(ta, "allowed", return_value=True):
            _run(ta.recv_ta_county(update, ctx))
        self.assertEqual(ctx.user_data["ta_session"].county, "")


class TestRecvTaPeriod(unittest.TestCase):
    def test_each_bucket_stores_correct_days_and_label(self):
        for label, days in ta._TA_PERIOD_OPTIONS:
            update = _make_query_update(f"ta_period:{days}")
            ctx = MagicMock()
            ctx.user_data = {}
            with patch.object(ta, "allowed", return_value=True):
                state = _run(ta.recv_ta_period(update, ctx))
            self.assertEqual(state, ta.TA.CONFIRM)
            sess = ctx.user_data["ta_session"]
            self.assertEqual(sess.period_days, days)
            self.assertEqual(sess.period_label, label)

    def test_today_stores_zero_days(self):
        update = _make_query_update("ta_period:0")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(ta, "allowed", return_value=True):
            _run(ta.recv_ta_period(update, ctx))
        self.assertEqual(ctx.user_data["ta_session"].period_days, 0)
        self.assertEqual(ctx.user_data["ta_session"].period_label, "Today")

    def test_cancel_ends_conversation(self):
        update = _make_query_update("ta_period_cancel")
        ctx = MagicMock()
        ctx.bot.send_message = AsyncMock()
        with patch.object(ta, "allowed", return_value=True):
            state = _run(ta.recv_ta_period(update, ctx))
        self.assertEqual(state, ta.ConversationHandler.END)


class TestRecvTaConfirm(unittest.TestCase):
    def _session(self, ctx):
        ctx.user_data = {"ta_session": ta.TASession(county="nairobi", period_days=7, period_label="Past Week")}

    def test_no_cancels(self):
        update = _make_query_update("ta_confirm:no")
        ctx = MagicMock()
        ctx.bot.send_message = AsyncMock()
        with patch.object(ta, "allowed", return_value=True):
            state = _run(ta.recv_ta_confirm(update, ctx))
        self.assertEqual(state, ta.ConversationHandler.END)

    def test_missing_tokens_ends_with_error(self):
        update = _make_query_update("ta_confirm:yes")
        ctx = MagicMock()
        ctx.bot.send_message = AsyncMock()
        self._session(ctx)
        with patch.object(ta, "allowed", return_value=True), \
             patch.object(ta, "get_valid_tokens", return_value=None):
            state = _run(ta.recv_ta_confirm(update, ctx))
        self.assertEqual(state, ta.ConversationHandler.END)
        text = update.callback_query.edit_message_text.call_args[0][0]
        self.assertIn("expired", text)

    def test_valid_tokens_kicks_off_background_run(self):
        update = _make_query_update("ta_confirm:yes")
        ctx = MagicMock()
        self._session(ctx)
        fake_tokens = MagicMock()
        with patch.object(ta, "allowed", return_value=True), \
             patch.object(ta, "get_valid_tokens", return_value=fake_tokens), \
             patch.object(ta.asyncio, "ensure_future") as mock_ensure, \
             patch.object(ta.asyncio, "to_thread") as mock_to_thread, \
             patch.object(ta.asyncio, "get_event_loop", return_value=MagicMock()):
            state = _run(ta.recv_ta_confirm(update, ctx))
        self.assertEqual(state, ta.ConversationHandler.END)
        mock_ensure.assert_called_once()
        mock_to_thread.assert_called_once()
        call_args = mock_to_thread.call_args[0]
        self.assertEqual(call_args[0], ta._ta_run)


if __name__ == "__main__":
    unittest.main()
