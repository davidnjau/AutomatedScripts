#!/usr/bin/env python3
"""
Unit tests for job_distribution.py — _jd_county_keyboard's toggle rendering,
_jd_fetch_task_detail's token-rotation/retry behavior, and the
cmd_job_distribution/recv_jd_cred/recv_jd_county conversation handlers.

Run with: python3 -m unittest discover -s assign/tests -v
"""

import asyncio
import os
import sys
import unittest
from unittest.mock import AsyncMock, MagicMock, patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import job_distribution as jd
import token_rotator
from ardhisasa_auth import AuthTokens
from token_rotator import _AllTokensExhausted, _TokenRotator

TOKENS = AuthTokens(access_token="acc", jwt="jwt")


def _load_workbook(xlsx_bytes):
    import io
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def _sheet_rows(ws):
    return [list(row) for row in ws.iter_rows(min_row=1, values_only=True)]


def _run(coro):
    return asyncio.run(coro)


def _make_update_with_message(text=""):
    update = MagicMock()
    update.message.text = text
    update.message.reply_text = AsyncMock()
    return update


def _make_update_with_callback(data):
    update = MagicMock()
    query = update.callback_query
    query.data = data
    query.answer = AsyncMock()
    query.edit_message_text = AsyncMock()
    query.edit_message_reply_markup = AsyncMock()
    query.message.reply_text = AsyncMock()
    query.message.chat_id = 555
    return update


class TestJdCountyKeyboard(unittest.TestCase):
    def test_selected_counties_are_checked(self):
        kbd = jd._jd_county_keyboard(["NAIROBI"])
        flat = [btn.text for row in kbd.inline_keyboard for btn in row]
        self.assertTrue(any(t.startswith("✅") and "Nairobi" in t for t in flat))
        self.assertTrue(any(t.startswith("☐") and "Kiambu" in t for t in flat))

    def test_includes_all_none_and_run_cancel_buttons(self):
        kbd = jd._jd_county_keyboard([])
        flat_data = [btn.callback_data for row in kbd.inline_keyboard for btn in row]
        self.assertIn("jd_county:ALL", flat_data)
        self.assertIn("jd_county:NONE", flat_data)
        self.assertIn("jd_county:done", flat_data)
        self.assertIn("jd_county:cancel", flat_data)


class TestJdFetchTaskDetail(unittest.TestCase):
    def test_success_returns_json(self):
        rotator = _TokenRotator([("staff", TOKENS)])
        fake_session = MagicMock()
        fake_session.get.return_value = MagicMock(
            status_code=200, raise_for_status=lambda: None, json=lambda: {"node": "X"}
        )
        result = jd._jd_fetch_task_detail(fake_session, rotator, "task-1")
        self.assertEqual(result, {"node": "X"})

    def test_all_tokens_exhausted_raises(self):
        rotator = _TokenRotator([])
        fake_session = MagicMock()
        with self.assertRaises(_AllTokensExhausted):
            jd._jd_fetch_task_detail(fake_session, rotator, "task-1")

    def test_403_rotates_then_succeeds(self):
        rotator = _TokenRotator([("staff", TOKENS), ("staff2", TOKENS)])
        fake_session = MagicMock()
        fake_session.get.side_effect = [
            MagicMock(status_code=403),
            MagicMock(status_code=200, raise_for_status=lambda: None, json=lambda: {"node": "Y"}),
        ]
        with patch.object(token_rotator.time, "sleep"):
            result = jd._jd_fetch_task_detail(fake_session, rotator, "task-1")
        self.assertEqual(result, {"node": "Y"})

    def test_exhausts_all_retries_on_persistent_5xx(self):
        rotator = _TokenRotator([("staff", TOKENS)])
        fake_session = MagicMock()
        fake_session.get.return_value = MagicMock(status_code=503)
        with patch.object(token_rotator.time, "sleep"):
            with self.assertRaises(RuntimeError):
                jd._jd_fetch_task_detail(fake_session, rotator, "task-1")


class TestJdFetchTeamMembers(unittest.TestCase):
    def test_paginates_until_no_next(self):
        fake_session = MagicMock()
        fake_session.get.side_effect = [
            MagicMock(raise_for_status=lambda: None, json=lambda: {"results": [{"userid": "1"}], "next": True}),
            MagicMock(raise_for_status=lambda: None, json=lambda: {"results": [{"userid": "2"}], "next": False}),
        ]
        members = jd._jd_fetch_team_members(fake_session, {}, "team-1")
        self.assertEqual(len(members), 2)
        self.assertEqual(fake_session.get.call_count, 2)


class TestCmdJobDistribution(unittest.TestCase):
    def test_no_valid_tokens_ends_conversation(self):
        update = _make_update_with_message()
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(jd, "allowed", return_value=True), \
             patch.object(jd, "_be_cred_keyboard", return_value=None):
            result = _run(jd.cmd_job_distribution(update, ctx))
        self.assertEqual(result, jd.ConversationHandler.END)

    def test_valid_tokens_ask_for_credential(self):
        update = _make_update_with_message()
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(jd, "allowed", return_value=True), \
             patch.object(jd, "_be_cred_keyboard", return_value=MagicMock()):
            result = _run(jd.cmd_job_distribution(update, ctx))
        self.assertEqual(result, jd.JD.PICK_CRED)


class TestRecvJdCred(unittest.TestCase):
    def test_resets_counties_to_all_and_asks_county(self):
        update = _make_update_with_callback("be_cred:staff")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession(counties=["NAIROBI"])}
        with patch.object(jd, "allowed", return_value=True):
            result = _run(jd.recv_jd_cred(update, ctx))
        self.assertEqual(result, jd.JD.COUNTY)
        self.assertEqual(ctx.user_data["jd_session"].counties, jd._JD_COUNTY_KEYS)


class TestJdBuildExcel(unittest.TestCase):
    def setUp(self):
        self.teams = [
            {"id": "t1", "team_name": "Team A", "min_amount": 0, "max_amount": 1000000},
            {"id": "t2", "team_name": "Team B", "min_amount": 0, "max_amount": 1000000},
        ]
        self.members_by_team = {
            "t1": [
                {"userid": "u1", "name": "Alice", "account_number": "A1",
                 "availability": "AVAILABLE", "registry": "R1"},
                {"userid": "u2", "name": "Bob", "account_number": "A2",
                 "availability": "AVAILABLE", "registry": "R1"},
            ],
            "t2": [
                {"userid": "u3", "name": "Carol", "account_number": "A3",
                 "availability": "AVAILABLE", "registry": "R2"},
            ],
        }
        # u1: 4 tasks, u2: 0 tasks, u3: 2 tasks → total assigned = 6
        self.tasks_by_userid = {
            "u1": [{"reference_number": f"REF-{i}"} for i in range(4)],
            "u3": [{"reference_number": f"REF-C{i}"} for i in range(2)],
        }
        self.unassigned_tasks = [
            {"reference_number": "PEND-1"}, {"reference_number": "PEND-2"},
        ]

    def _build(self):
        xlsx_bytes = jd._jd_build_excel(
            self.teams, self.members_by_team, self.tasks_by_userid, self.unassigned_tasks
        )
        return _load_workbook(xlsx_bytes)

    def test_team_summary_has_pending_tasks_column_on_every_row(self):
        wb   = self._build()
        rows = _sheet_rows(wb["Team Summary"])
        header = rows[0]
        self.assertIn("Pending Tasks", header)
        pending_col = header.index("Pending Tasks")
        assigned_col = header.index("Assigned Tasks")

        # Team A row and Team B row must both show the true pending total (2), not blank.
        team_rows = [r for r in rows[1:] if r[0] in ("Team A", "Team B")]
        self.assertEqual(len(team_rows), 2)
        for r in team_rows:
            self.assertEqual(r[pending_col], 2)

        # Team A: assigned=4, Team B: assigned=2
        by_name = {r[0]: r for r in team_rows}
        self.assertEqual(by_name["Team A"][assigned_col], 4)
        self.assertEqual(by_name["Team B"][assigned_col], 2)

    def test_team_summary_has_grand_total_row(self):
        wb   = self._build()
        rows = _sheet_rows(wb["Team Summary"])
        header = rows[0]
        pending_col  = header.index("Pending Tasks")
        assigned_col = header.index("Assigned Tasks")

        total_row = next(r for r in rows if r[0] == "TOTAL")
        self.assertEqual(total_row[assigned_col], 6)
        self.assertEqual(total_row[pending_col], 2)

    def test_member_distribution_lists_below_average_members(self):
        wb   = self._build()
        rows = _sheet_rows(wb["Member Distribution"])
        first_cells = [r[0] for r in rows]

        overall_idx = next(
            i for i, v in enumerate(first_cells)
            if isinstance(v, str) and v.startswith("MEMBERS BELOW OVERALL AVERAGE")
        )
        team_idx = next(
            i for i, v in enumerate(first_cells)
            if v == "MEMBERS BELOW THEIR OWN TEAM'S AVERAGE"
        )

        # overall avg = 6 tasks / 3 members = 2.0 → only Bob (0 tasks) is below it
        overall_section = rows[overall_idx:team_idx]
        self.assertIn("Bob", [r[0] for r in overall_section])
        self.assertNotIn("Alice", [r[0] for r in overall_section])
        self.assertNotIn("Carol", [r[0] for r in overall_section])

        # Team A avg = 4/2 = 2.0, Team B avg = 2/1 = 2.0 → only Bob (0) is below its team avg
        team_section = rows[team_idx:]
        self.assertIn("Bob", [r[0] for r in team_section])
        self.assertNotIn("Alice", [r[0] for r in team_section])
        self.assertNotIn("Carol", [r[0] for r in team_section])


class TestRecvJdCounty(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_update_with_callback("jd_county:cancel")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession()}
        ctx.bot.send_message = AsyncMock()
        with patch.object(jd, "allowed", return_value=True):
            result = _run(jd.recv_jd_county(update, ctx))
        self.assertEqual(result, jd.ConversationHandler.END)

    def test_done_with_no_counties_reprompts(self):
        update = _make_update_with_callback("jd_county:done")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession(counties=[])}
        with patch.object(jd, "allowed", return_value=True):
            result = _run(jd.recv_jd_county(update, ctx))
        self.assertEqual(result, jd.JD.COUNTY)

    def test_done_with_expired_tokens_ends_conversation(self):
        update = _make_update_with_callback("jd_county:done")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession(cred_type="staff")}
        ctx.bot.send_message = AsyncMock()
        with patch.object(jd, "allowed", return_value=True), \
             patch.object(jd, "get_valid_tokens", return_value=None):
            result = _run(jd.recv_jd_county(update, ctx))
        self.assertEqual(result, jd.ConversationHandler.END)

    def test_done_with_valid_tokens_schedules_background_run(self):
        update = _make_update_with_callback("jd_county:done")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession(cred_type="staff")}
        ctx.bot.send_message = AsyncMock()
        with patch.object(jd, "allowed", return_value=True), \
             patch.object(jd, "get_valid_tokens", return_value=TOKENS), \
             patch.object(jd.asyncio, "to_thread", new=AsyncMock()), \
             patch.object(jd.asyncio, "ensure_future", side_effect=lambda coro: coro.close()):
            result = _run(jd.recv_jd_county(update, ctx))
        self.assertEqual(result, jd.ConversationHandler.END)

    def test_toggle_all(self):
        update = _make_update_with_callback("jd_county:ALL")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession(counties=[])}
        with patch.object(jd, "allowed", return_value=True):
            result = _run(jd.recv_jd_county(update, ctx))
        self.assertEqual(result, jd.JD.COUNTY)
        self.assertEqual(ctx.user_data["jd_session"].counties, jd._JD_COUNTY_KEYS)

    def test_toggle_none(self):
        update = _make_update_with_callback("jd_county:NONE")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession()}
        with patch.object(jd, "allowed", return_value=True):
            result = _run(jd.recv_jd_county(update, ctx))
        self.assertEqual(ctx.user_data["jd_session"].counties, [])

    def test_toggle_single_county_add_and_remove(self):
        update = _make_update_with_callback("jd_county:NAIROBI")
        ctx = MagicMock()
        ctx.user_data = {"jd_session": jd.JDSession(counties=[])}
        with patch.object(jd, "allowed", return_value=True):
            _run(jd.recv_jd_county(update, ctx))
        self.assertIn("NAIROBI", ctx.user_data["jd_session"].counties)

        with patch.object(jd, "allowed", return_value=True):
            _run(jd.recv_jd_county(update, ctx))
        self.assertNotIn("NAIROBI", ctx.user_data["jd_session"].counties)


if __name__ == "__main__":
    unittest.main()
