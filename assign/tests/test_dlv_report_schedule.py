#!/usr/bin/env python3
"""
Unit tests for dlv_report_schedule.py — schedule persistence, the Excel
export builder, the background job, and the Add-schedule conversation
handlers (scope/valuer/tag/period/interval/email).

Run with: python3 -m unittest discover -s assign/tests -v
"""

import asyncio
import os
import sys
import tempfile
import unittest
from unittest.mock import AsyncMock, MagicMock, patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import dlv_report_schedule as drs


def _run(coro):
    return asyncio.run(coro)


def _make_query_update(data):
    update = MagicMock()
    query = update.callback_query
    query.data = data
    query.answer = AsyncMock()
    query.edit_message_text = AsyncMock()
    query.message.reply_text = AsyncMock()
    return update


def _make_message_update(text):
    update = MagicMock()
    update.message.text = text
    update.message.reply_text = AsyncMock()
    return update


class TestSchedulePersistence(unittest.TestCase):
    """load/add/get/remove_dlv_report_schedule(s) — the persistence layer."""

    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.sched_file = os.path.join(self.tmpdir.name, "saved_dlv_report_schedules.json")
        self._patch = patch.object(drs, "SAVED_DRS_FILE", self.sched_file)
        self._patch.start()

    def tearDown(self):
        self._patch.stop()
        self.tmpdir.cleanup()

    def test_load_missing_file_returns_empty_list(self):
        self.assertEqual(drs.load_dlv_report_schedules(), [])

    def test_add_then_load_roundtrip(self):
        schedule_id = drs.add_dlv_report_schedule({"scope": "valuer", "email": "a@b.com"})
        schedules = drs.load_dlv_report_schedules()
        self.assertEqual(len(schedules), 1)
        self.assertEqual(schedules[0]["email"], "a@b.com")
        self.assertEqual(schedules[0]["id"], schedule_id)

    def test_add_multiple_schedules_get_distinct_ids(self):
        id1 = drs.add_dlv_report_schedule({"email": "a@b.com"})
        id2 = drs.add_dlv_report_schedule({"email": "c@d.com"})
        self.assertNotEqual(id1, id2)
        self.assertEqual(len(drs.load_dlv_report_schedules()), 2)

    def test_get_dlv_report_schedule_finds_by_id(self):
        schedule_id = drs.add_dlv_report_schedule({"email": "a@b.com"})
        cfg = drs.get_dlv_report_schedule(schedule_id)
        self.assertEqual(cfg["email"], "a@b.com")

    def test_get_dlv_report_schedule_returns_none_when_missing(self):
        self.assertIsNone(drs.get_dlv_report_schedule("no-such-id"))

    def test_remove_dlv_report_schedule_deletes_only_that_one(self):
        id1 = drs.add_dlv_report_schedule({"email": "a@b.com"})
        id2 = drs.add_dlv_report_schedule({"email": "c@d.com"})
        removed = drs.remove_dlv_report_schedule(id1)
        self.assertTrue(removed)
        remaining = drs.load_dlv_report_schedules()
        self.assertEqual(len(remaining), 1)
        self.assertEqual(remaining[0]["id"], id2)

    def test_remove_missing_schedule_returns_false(self):
        self.assertFalse(drs.remove_dlv_report_schedule("no-such-id"))


class TestFormatScheduleSummary(unittest.TestCase):
    """_drs_format_schedule_summary — the one-line schedule description."""

    def test_valuer_scope_summary(self):
        summary = drs._drs_format_schedule_summary({
            "scope": "valuer", "target_name": "Jane Doe", "period_label": "1 week",
            "interval_minutes": 60, "email": "a@b.com",
        })
        self.assertIn("👤 Valuer", summary)
        self.assertIn("Jane Doe", summary)
        self.assertIn("1 week", summary)
        self.assertIn("60 min", summary)
        self.assertIn("a@b.com", summary)

    def test_tag_scope_summary(self):
        summary = drs._drs_format_schedule_summary({
            "scope": "tag", "target_name": "Queue", "period_label": "All time",
            "interval_minutes": 1440, "email": "a@b.com",
        })
        self.assertIn("🏷 Tag", summary)
        self.assertIn("Queue", summary)

    def test_special_chars_in_name_and_email_are_escaped(self):
        summary = drs._drs_format_schedule_summary({
            "scope": "valuer", "target_name": "Jane_Doe", "period_label": "All time",
            "interval_minutes": 60, "email": "john_doe@example.com",
        })
        self.assertIn("Jane\\_Doe", summary)
        self.assertIn("john\\_doe@example.com", summary)


class TestBuildExcel(unittest.TestCase):
    """_drs_build_excel — one sheet per section, populated from item dicts."""

    def test_three_sheets_created_with_expected_titles(self):
        xlsx_bytes = drs._drs_build_excel([], [], [])
        import io
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertEqual(wb.sheetnames, ["Currently Queued", "At Valuer's Desk", "Valuer Completed"])

    def test_rows_populated_per_section_with_own_date_field(self):
        import io
        queued = [{"ref": "REF1", "valuer_name": "Jane Doe", "assessor": "A1", "consideration": "1000000",
                   "currency_code": "KES", "parcel": "P1", "tag": "Queue", "queued_at": "2026-07-10T10:00:00"}]
        desk = [{"ref": "REF2", "assigned_at": "2026-07-15 09:00:00"}]
        closed = [{"ref": "REF3", "closed_at": "2026-07-12T11:00:00"}]
        xlsx_bytes = drs._drs_build_excel(queued, desk, closed)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

        queued_ws = wb["Currently Queued"]
        self.assertEqual(queued_ws.cell(row=2, column=1).value, "REF1")
        self.assertEqual(queued_ws.cell(row=2, column=2).value, "Jane Doe")
        self.assertEqual(queued_ws.cell(row=2, column=8).value, "2026-07-10T10:00:00")

        desk_ws = wb["At Valuer's Desk"]
        self.assertEqual(desk_ws.cell(row=2, column=1).value, "REF2")
        self.assertEqual(desk_ws.cell(row=2, column=8).value, "2026-07-15 09:00:00")

        closed_ws = wb["Valuer Completed"]
        self.assertEqual(closed_ws.cell(row=2, column=1).value, "REF3")
        self.assertEqual(closed_ws.cell(row=2, column=8).value, "2026-07-12T11:00:00")


class TestDrsJob(unittest.TestCase):
    """_drs_job — the repeating job that rebuilds and emails one schedule."""

    def test_schedule_not_found_removes_orphaned_job(self):
        ctx = MagicMock()
        ctx.job.data = "missing-id"
        with patch.object(drs, "get_dlv_report_schedule", return_value=None):
            _run(drs._drs_job(ctx))
        ctx.job.schedule_removal.assert_called_once()

    def test_successful_run_sends_email(self):
        ctx = MagicMock()
        ctx.job.data = "sched-1"
        cfg = {"scope": "valuer", "target_key": "u1", "target_name": "Jane Doe",
               "period_days": 7, "email": "a@b.com"}
        with patch.object(drs, "get_dlv_report_schedule", return_value=cfg), \
             patch.object(drs, "_dt_gather_report_data", return_value=([{"ref": "REF1"}], [], [])), \
             patch.object(drs, "_send_bulk_export_email") as mock_send:
            _run(drs._drs_job(ctx))
        mock_send.assert_called_once()
        self.assertEqual(mock_send.call_args[0][0], "a@b.com")

    def test_email_failure_is_logged_not_raised(self):
        ctx = MagicMock()
        ctx.job.data = "sched-1"
        cfg = {"scope": "tag", "target_key": "Queue", "target_name": "Queue",
               "period_days": 0, "email": "a@b.com"}
        with patch.object(drs, "get_dlv_report_schedule", return_value=cfg), \
             patch.object(drs, "_dt_gather_report_data", return_value=([], [], [])), \
             patch.object(drs, "_send_bulk_export_email", side_effect=RuntimeError("smtp down")):
            _run(drs._drs_job(ctx))   # must not raise


class TestCmdDrs(unittest.TestCase):
    """cmd_drs — entry point, lists existing schedules."""

    def test_no_schedules_shows_empty_message(self):
        update = _make_message_update("")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "load_dlv_report_schedules", return_value=[]):
            result = _run(drs.cmd_drs(update, ctx))
        self.assertEqual(result, drs.DRS.MENU)
        text = update.message.reply_text.call_args[0][0]
        self.assertIn("No active schedules", text)

    def test_existing_schedules_are_listed(self):
        update = _make_message_update("")
        ctx = MagicMock()
        ctx.user_data = {}
        cfg = {"scope": "valuer", "target_name": "Jane Doe", "period_label": "All time",
               "interval_minutes": 60, "email": "a@b.com"}
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "load_dlv_report_schedules", return_value=[cfg]):
            _run(drs.cmd_drs(update, ctx))
        text = update.message.reply_text.call_args[0][0]
        self.assertIn("1 active schedule", text)
        self.assertIn("Jane Doe", text)


class TestRecvDrsMenu(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_menu:cancel")
        ctx = MagicMock()
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_menu(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_add_shows_scope_picker(self):
        update = _make_query_update("drs_menu:add")
        ctx = MagicMock()
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_menu(update, ctx))
        self.assertEqual(result, drs.DRS.SCOPE)

    def test_remove_shows_remove_picker(self):
        update = _make_query_update("drs_menu:remove")
        ctx = MagicMock()
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "load_dlv_report_schedules", return_value=[{"id": "s1", "email": "a@b.com"}]):
            result = _run(drs.recv_drs_menu(update, ctx))
        self.assertEqual(result, drs.DRS.REMOVE_PICK)


class TestRecvDrsRemove(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_remove:cancel")
        ctx = MagicMock()
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_remove(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_removes_schedule_and_cancels_its_job(self):
        update = _make_query_update("drs_remove:sched-1")
        ctx = MagicMock()
        job = MagicMock()
        ctx.job_queue.get_jobs_by_name.return_value = [job]
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "remove_dlv_report_schedule", return_value=True) as mock_remove:
            result = _run(drs.recv_drs_remove(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)
        mock_remove.assert_called_once_with("sched-1")
        job.schedule_removal.assert_called_once()
        ctx.job_queue.get_jobs_by_name.assert_called_once_with("drs_job:sched-1")


class TestRecvDrsScope(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_scope:cancel")
        ctx = MagicMock()
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_scope(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_tag_scope_shows_tag_picker(self):
        update = _make_query_update("drs_scope:tag")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_scope(update, ctx))
        self.assertEqual(result, drs.DRS.PICK_TAG)
        self.assertEqual(drs._get_drs_sess(ctx).scope, "tag")

    def test_valuer_scope_with_no_valuers_ends_conversation(self):
        update = _make_query_update("drs_scope:valuer")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "_dt_collect_valuers", return_value=[]):
            result = _run(drs.recv_drs_scope(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_valuer_scope_shows_valuer_picker(self):
        update = _make_query_update("drs_scope:valuer")
        ctx = MagicMock()
        ctx.user_data = {}
        valuers = [{"key": "u1", "name": "Jane Doe"}]
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "_dt_collect_valuers", return_value=valuers):
            result = _run(drs.recv_drs_scope(update, ctx))
        self.assertEqual(result, drs.DRS.PICK_VALUER)
        self.assertEqual(drs._get_drs_sess(ctx).valuer_choices, valuers)


class TestRecvDrsPickValuer(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_pickvaluer:cancel")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_pick_valuer(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_valid_index_stores_target_and_moves_to_period(self):
        update = _make_query_update("drs_pickvaluer:0")
        ctx = MagicMock()
        ctx.user_data = {}
        sess = drs._get_drs_sess(ctx)
        sess.valuer_choices = [{"key": "u1", "name": "Jane_Doe"}]
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_pick_valuer(update, ctx))
        self.assertEqual(result, drs.DRS.PERIOD)
        self.assertEqual(sess.target_key, "u1")
        self.assertEqual(sess.target_name, "Jane_Doe")
        text = update.callback_query.edit_message_text.call_args[0][0]
        self.assertIn("Jane\\_Doe", text)

    def test_out_of_range_index_stays_on_picker(self):
        update = _make_query_update("drs_pickvaluer:5")
        ctx = MagicMock()
        ctx.user_data = {}
        sess = drs._get_drs_sess(ctx)
        sess.valuer_choices = [{"key": "u1", "name": "Jane Doe"}]
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_pick_valuer(update, ctx))
        self.assertEqual(result, drs.DRS.PICK_VALUER)


class TestRecvDrsPickTag(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_picktag:cancel")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_pick_tag(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_picking_a_tag_stores_target_and_moves_to_period(self):
        update = _make_query_update("drs_picktag:Queue")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_pick_tag(update, ctx))
        self.assertEqual(result, drs.DRS.PERIOD)
        sess = drs._get_drs_sess(ctx)
        self.assertEqual(sess.target_key, "Queue")
        self.assertEqual(sess.target_name, "Queue")


class TestRecvDrsPeriod(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_period:cancel")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_period(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_picking_a_period_moves_to_interval(self):
        update = _make_query_update("drs_period:7")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_period(update, ctx))
        self.assertEqual(result, drs.DRS.INTERVAL)
        sess = drs._get_drs_sess(ctx)
        self.assertEqual(sess.period_days, 7)
        self.assertEqual(sess.period_label, "1 week")


class TestRecvDrsInterval(unittest.TestCase):
    def test_cancel_ends_conversation(self):
        update = _make_query_update("drs_interval:cancel")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_interval(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)

    def test_picking_an_interval_moves_to_email(self):
        update = _make_query_update("drs_interval:120")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_interval(update, ctx))
        self.assertEqual(result, drs.DRS.EMAIL)
        self.assertEqual(drs._get_drs_sess(ctx).interval_minutes, 120)


class TestRecvDrsEmail(unittest.TestCase):
    def test_invalid_email_reprompts(self):
        update = _make_message_update("not-an-email")
        ctx = MagicMock()
        ctx.user_data = {}
        with patch.object(drs, "allowed", return_value=True):
            result = _run(drs.recv_drs_email(update, ctx))
        self.assertEqual(result, drs.DRS.EMAIL)

    def test_valid_email_saves_schedule_and_registers_job(self):
        update = _make_message_update("a@b.com")
        ctx = MagicMock()
        ctx.user_data = {}
        sess = drs._get_drs_sess(ctx)
        sess.scope = "valuer"
        sess.target_key = "u1"
        sess.target_name = "Jane Doe"
        sess.period_days = 7
        sess.period_label = "1 week"
        sess.interval_minutes = 60
        with patch.object(drs, "allowed", return_value=True), \
             patch.object(drs, "add_dlv_report_schedule", return_value="sched-1") as mock_add:
            result = _run(drs.recv_drs_email(update, ctx))
        self.assertEqual(result, drs.ConversationHandler.END)
        mock_add.assert_called_once()
        saved_cfg = mock_add.call_args[0][0]
        self.assertEqual(saved_cfg["email"], "a@b.com")
        self.assertEqual(saved_cfg["target_key"], "u1")
        ctx.job_queue.run_repeating.assert_called_once()
        self.assertEqual(ctx.job_queue.run_repeating.call_args.kwargs["name"], "drs_job:sched-1")
        self.assertEqual(ctx.job_queue.run_repeating.call_args.kwargs["interval"], 3600)


if __name__ == "__main__":
    unittest.main()
