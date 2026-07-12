#!/usr/bin/env python3
"""
excel_report.py
================
Shared openpyxl styling boilerplate for the Excel report builders in this
bot (DLV Tasks, Bulk Export, Job Distribution, Valuer Tasks) — each
independently repeated the same bold-blue-header/autofilter/frozen-pane/
auto-width pattern before this was pulled out.

Feature-specific styling (row banding, section headers, multi-sheet
layouts, number/date formats) stays in each feature module — this only
covers the "style row 1 as a header and autofilter/freeze it" and
"auto-size every column" steps that were identical everywhere.
"""

from typing import List, Optional

from openpyxl.styles import Font, PatternFill

_DEFAULT_HEADER_FONT = Font(bold=True)
_DEFAULT_HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")


def style_header_row(
    ws,
    columns: List[str],
    font: Optional[Font] = None,
    fill: Optional[PatternFill] = None,
) -> None:
    """
    Append `columns` as row 1, bold+fill each header cell, and turn on
    autofilter + a frozen top row. Matches the header treatment every
    Excel report in this bot already used.
    """
    ws.append(columns)
    header_font = font or _DEFAULT_HEADER_FONT
    header_fill = fill or _DEFAULT_HEADER_FILL
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes    = "A2"


def autofit_columns(ws, min_width: int = 12, max_width: int = 55) -> None:
    """
    Auto-size every column in `ws` to fit its longest cell value (header
    included), clamped to [min_width, max_width].
    """
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=min_width)
        ws.column_dimensions[col[0].column_letter].width = max(min_width, min(max_width, max_len + 2))
