#!/usr/bin/env python3
"""
dlv_report_schedule.py
=======================
DLV Report Schedule — recurring, emailed DLV Tasks By Valuer/By Tag
reports (📧 DLV Report Schedule button / /dlvreportschedule). This is an
additional delivery option alongside DLV Tasks' own interactive
Telegram-only By Valuer/By Tag flow — it does not replace it.

Each schedule is a dict with its own generated "id", stored as a list in
saved_dlv_report_schedules.json, and gets its own repeating PTB job —
named f"drs_job:{schedule_id}" and tagged with that id via the job's
`data` param — mirroring auto_fetch.py's per-schedule job pattern.

The report itself is entirely sourced from local JSON stores (no live
API calls, same as DLV Tasks' own By Valuer/By Tag flow), so no
credential/token step is needed anywhere in this flow — it reuses
dlv_tasks.py's _dt_gather_report_data (the exact same Currently
Queued/At Valuer's Desk/Valuer Completed filtering logic that report
already uses) and _dt_collect_valuers (the valuer picker's data source),
and emails the result as an Excel workbook (one sheet per section) via
email_service._send_bulk_export_email — the same attachment-based sender
DLV Tasks' own Open Tasks email option already uses.

Call register(app) from bot.py's main() to wire this feature in (this
also restores one repeating job per saved schedule on startup).
"""

import io
import json
import os
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum, auto
from typing import Dict, List, Optional

import openpyxl
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

from common import (
    BTN_DLV_REPORT_SCHEDULE,
    DATA_DIR,
    _atomic_json_write,
    _CANCEL_FILTER,
    _main_menu,
    allowed,
    cmd_cancel,
    deny,
    fallback,
    logger,
    md_escape,
    not_cancel,
)
from dlv_core import DLV_TAGS
from dlv_tasks import _DT_PERIOD_OPTIONS, _dt_collect_valuers, _dt_gather_report_data
from email_service import _send_bulk_export_email
from excel_report import autofit_columns, style_header_row

SAVED_DRS_FILE = os.path.join(DATA_DIR, "saved_dlv_report_schedules.json")

_DRS_EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[a-zA-Z]{2,}")

# Same interval choices as Auto Fetch's schedule picker (auto_fetch.py's
# _AF_INTERVALS) — duplicated rather than shared since it's a small literal
# list local to each feature's own picker keyboard.
_DRS_INTERVALS = [
    ("1 hr",    60),
    ("2 hr",   120),
    ("4 hr",   240),
    ("6 hr",   360),
    ("12 hr",  720),
    ("24 hr", 1440),
]

# Mirrors auto_fetch.py's _AF_RESTORE_FIRST_RUN_DELAY reasoning: run soon
# after startup rather than waiting a full interval, so a restart doesn't
# push the first post-restart run out by up to the whole interval.
_DRS_RESTORE_FIRST_RUN_DELAY = 60   # seconds


# ──────────────────────────────────────────────────────────
# States — DLV Report Schedule conversation
# ──────────────────────────────────────────────────────────
class DRS(Enum):
    MENU        = auto()   # list existing schedules; choose Add / Remove
    REMOVE_PICK = auto()   # pick which schedule to remove
    SCOPE       = auto()   # By Valuer or By Tag
    PICK_VALUER = auto()   # valuer picker (scope == "valuer")
    PICK_TAG    = auto()   # tag picker (scope == "tag")
    PERIOD      = auto()   # look-back period picker
    INTERVAL    = auto()   # how often to run
    EMAIL       = auto()   # recipient email address


# ──────────────────────────────────────────────────────────
# Per-user session data
# ──────────────────────────────────────────────────────────
@dataclass
class DRSSession:
    scope:            str        = ""    # "valuer" | "tag"
    valuer_choices:   List[Dict] = field(default_factory=list)   # By Valuer picker: [{key, name}]
    target_key:       str        = ""    # valuer key (uid/name) or tag string
    target_name:      str        = ""    # display name — same as target_key for tags
    period_days:      int        = 0
    period_label:     str        = "All time"
    interval_minutes: int        = 60


def _get_drs_sess(ctx: ContextTypes.DEFAULT_TYPE) -> DRSSession:
    """Fetch (creating if absent) this user's in-progress schedule-setup session."""
    if "drs_session" not in ctx.user_data:
        ctx.user_data["drs_session"] = DRSSession()
    return ctx.user_data["drs_session"]


# ──────────────────────────────────────────────────────────
# Persistence — saved_dlv_report_schedules.json
# ──────────────────────────────────────────────────────────
def load_dlv_report_schedules() -> List[Dict]:
    """Load every saved DLV Report schedule."""
    try:
        with open(SAVED_DRS_FILE) as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def save_dlv_report_schedules(schedules: List[Dict]) -> None:
    """Persist the full list of DLV Report schedules."""
    _atomic_json_write(SAVED_DRS_FILE, schedules, indent=2)


def get_dlv_report_schedule(schedule_id: str) -> Optional[Dict]:
    """Look up a single saved schedule by id."""
    return next((s for s in load_dlv_report_schedules() if s.get("id") == schedule_id), None)


def add_dlv_report_schedule(cfg: Dict) -> str:
    """Assign a new id to cfg, append it to the saved list, and return the id."""
    schedule_id = str(uuid.uuid4())[:8]
    schedules = load_dlv_report_schedules()
    schedules.append({**cfg, "id": schedule_id})
    save_dlv_report_schedules(schedules)
    return schedule_id


def remove_dlv_report_schedule(schedule_id: str) -> bool:
    """Delete a schedule by id. Returns False if no schedule had that id."""
    schedules = load_dlv_report_schedules()
    remaining = [s for s in schedules if s.get("id") != schedule_id]
    if len(remaining) == len(schedules):
        return False
    save_dlv_report_schedules(remaining)
    return True


def _drs_format_schedule_summary(cfg: Dict) -> str:
    """One-line summary of a schedule's settings, used in the schedule list and remove picker."""
    scope_label = "🏷 Tag" if cfg.get("scope") == "tag" else "👤 Valuer"
    return (
        f"{scope_label}: {md_escape(cfg.get('target_name', '?'))} | "
        f"{cfg.get('period_label', 'All time')} | every {cfg.get('interval_minutes', '?')} min | "
        f"{md_escape(cfg.get('email', '?'))}"
    )


# ──────────────────────────────────────────────────────────
# Excel export — one sheet per report section
# ──────────────────────────────────────────────────────────
def _drs_populate_sheet(ws, items: List[Dict], date_field: str) -> None:
    """Fill one sheet with a section's items: Reference, Valuer, Assessor,
    Consideration, Currency, Parcel, Tag, and the section's own date field
    (queued_at/assigned_at/closed_at)."""
    cols = ["Reference Number", "Valuer", "Assessor", "Consideration", "Currency", "Parcel", "Tag", "Date"]
    style_header_row(ws, cols)
    for item in items:
        ws.append([
            item.get("ref", ""),
            item.get("valuer_name", ""),
            item.get("assessor", ""),
            item.get("consideration_amount") or item.get("consideration") or "",
            item.get("currency_code", ""),
            item.get("parcel") or item.get("parcel_number") or "",
            item.get("tag", ""),
            item.get(date_field, ""),
        ])
    autofit_columns(ws, min_width=15, max_width=60)


def _drs_build_excel(queued: List[Dict], desk: List[Dict], closed: List[Dict]) -> bytes:
    """Build the DLV Report Schedule Excel export — one sheet per section
    (Currently Queued / At Valuer's Desk / Valuer Completed), each row a
    ref with its assessor/consideration/parcel/tag/date."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Currently Queued"
    _drs_populate_sheet(ws1, queued, "queued_at")
    ws2 = wb.create_sheet("At Valuer's Desk")
    _drs_populate_sheet(ws2, desk, "assigned_at")
    ws3 = wb.create_sheet("Valuer Completed")
    _drs_populate_sheet(ws3, closed, "closed_at")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ──────────────────────────────────────────────────────────
# Keyboards
# ──────────────────────────────────────────────────────────
def _drs_scope_keyboard() -> InlineKeyboardMarkup:
    """Choose between By Valuer and By Tag as the schedule's report scope."""
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("👤 By Valuer", callback_data="drs_scope:valuer")],
        [InlineKeyboardButton("🏷 By Tag",   callback_data="drs_scope:tag")],
        [InlineKeyboardButton("🛑 Cancel",   callback_data="drs_scope:cancel")],
    ])


def _drs_valuer_keyboard(valuers: List[Dict]) -> InlineKeyboardMarkup:
    """Two-per-row picker of valuer names, each tap selecting that valuer by index."""
    rows = []
    for i in range(0, len(valuers), 2):
        pair = valuers[i:i + 2]
        rows.append([
            InlineKeyboardButton(v["name"], callback_data=f"drs_pickvaluer:{i + j}")
            for j, v in enumerate(pair)
        ])
    rows.append([InlineKeyboardButton("🛑 Cancel", callback_data="drs_pickvaluer:cancel")])
    return InlineKeyboardMarkup(rows)


def _drs_tag_keyboard() -> InlineKeyboardMarkup:
    """Fixed-list tag picker, mirroring dlv_tasks.py's own By Tag picker."""
    rows = [[InlineKeyboardButton(t, callback_data=f"drs_picktag:{t}")] for t in DLV_TAGS]
    rows.append([InlineKeyboardButton("🛑 Cancel", callback_data="drs_picktag:cancel")])
    return InlineKeyboardMarkup(rows)


def _drs_period_keyboard() -> InlineKeyboardMarkup:
    """Look-back period picker for the Valuer Completed section, same options as DLV Tasks' own."""
    row = [InlineKeyboardButton(label, callback_data=f"drs_period:{days}") for label, days in _DT_PERIOD_OPTIONS]
    return InlineKeyboardMarkup([row, [InlineKeyboardButton("🛑 Cancel", callback_data="drs_period:cancel")]])


def _drs_interval_keyboard() -> InlineKeyboardMarkup:
    """How often the schedule should re-run and re-email the report."""
    rows = [[InlineKeyboardButton(label, callback_data=f"drs_interval:{mins}")] for label, mins in _DRS_INTERVALS]
    rows.append([InlineKeyboardButton("🛑 Cancel", callback_data="drs_interval:cancel")])
    return InlineKeyboardMarkup(rows)


def _drs_remove_keyboard(schedules: List[Dict]) -> InlineKeyboardMarkup:
    """One button per saved schedule, for the remove picker."""
    rows = [
        [InlineKeyboardButton(_drs_format_schedule_summary(cfg)[:60], callback_data=f"drs_remove:{cfg['id']}")]
        for cfg in schedules
    ]
    rows.append([InlineKeyboardButton("🛑 Cancel", callback_data="drs_remove:cancel")])
    return InlineKeyboardMarkup(rows)


# ──────────────────────────────────────────────────────────
# Background job — rebuild + email one saved schedule
# ──────────────────────────────────────────────────────────
async def _drs_job(context: ContextTypes.DEFAULT_TYPE) -> None:
    """PTB repeating job: rebuild and email one saved DLV Report schedule's
    Excel export. No live API calls or auth needed — everything's sourced
    from local JSON stores, same as DLV Tasks' own By Valuer/By Tag flow."""
    schedule_id = context.job.data
    cfg = get_dlv_report_schedule(schedule_id)
    if not cfg:
        logger.warning("DLV Report schedule %s not found — removing orphaned job.", schedule_id)
        context.job.schedule_removal()
        return

    queued, desk, closed = _dt_gather_report_data(cfg["scope"], cfg["target_key"], cfg["period_days"])
    xlsx_bytes = _drs_build_excel(queued, desk, closed)
    safe_name  = re.sub(r"[^A-Za-z0-9_-]+", "_", cfg.get("target_name", "report"))
    filename   = f"dlv_report_{safe_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    try:
        _send_bulk_export_email(cfg["email"], filename, xlsx_bytes)
        logger.info(
            "DLV Report schedule %s emailed to %s (%d queued / %d at-desk / %d completed).",
            schedule_id, cfg["email"], len(queued), len(desk), len(closed),
        )
    except Exception as e:
        logger.error("DLV Report schedule %s email failed: %s", schedule_id, e)


# ──────────────────────────────────────────────────────────
# Conversation handlers
# ──────────────────────────────────────────────────────────
async def cmd_drs(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Entry point (/dlvreportschedule or the menu button) — list existing
    schedules and offer Add/Remove."""
    if not allowed(update): return await deny(update)
    ctx.user_data["drs_session"] = DRSSession()
    schedules = load_dlv_report_schedules()

    lines = ["📧 *DLV Report Schedule*\n"]
    if schedules:
        lines.append(f"{len(schedules)} active schedule(s):")
        lines += [f"• {_drs_format_schedule_summary(cfg)}" for cfg in schedules]
    else:
        lines.append("No active schedules yet.")

    rows = [[InlineKeyboardButton("➕ Add Schedule", callback_data="drs_menu:add")]]
    if schedules:
        rows.append([InlineKeyboardButton("🗑 Remove", callback_data="drs_menu:remove")])
    rows.append([InlineKeyboardButton("🛑 Cancel", callback_data="drs_menu:cancel")])

    await update.message.reply_text("\n".join(lines), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(rows))
    return DRS.MENU


async def recv_drs_menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Handle the Add/Remove/Cancel choice from cmd_drs."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()
    action = query.data.split(":", 1)[1]

    if action == "cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    if action == "remove":
        schedules = load_dlv_report_schedules()
        await query.edit_message_text("🗑 Pick a schedule to remove:", reply_markup=_drs_remove_keyboard(schedules))
        return DRS.REMOVE_PICK

    await query.edit_message_text("Choose report scope:", reply_markup=_drs_scope_keyboard())
    return DRS.SCOPE


async def recv_drs_remove(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Delete the picked schedule and cancel its repeating job."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()
    schedule_id = query.data.split(":", 1)[1]

    if schedule_id == "cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    for job in ctx.job_queue.get_jobs_by_name(f"drs_job:{schedule_id}"):
        job.schedule_removal()
    removed = remove_dlv_report_schedule(schedule_id)

    await query.edit_message_text("🗑 Schedule removed." if removed else "⚠️ That schedule was already removed.")
    await query.message.reply_text("Main menu:", reply_markup=_main_menu())
    return ConversationHandler.END


async def recv_drs_scope(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Handle the By Valuer/By Tag scope choice, then show the matching picker."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()
    scope = query.data.split(":", 1)[1]

    if scope == "cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    sess = _get_drs_sess(ctx)
    sess.scope = scope

    if scope == "tag":
        await query.edit_message_text("🏷 Pick a tag:", reply_markup=_drs_tag_keyboard())
        return DRS.PICK_TAG

    valuers = _dt_collect_valuers()
    if not valuers:
        await query.edit_message_text("ℹ️ No queued, at-desk, or closed DLV tasks yet — nothing to schedule.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END
    sess.valuer_choices = valuers
    await query.edit_message_text("👤 Pick a valuer:", reply_markup=_drs_valuer_keyboard(valuers))
    return DRS.PICK_VALUER


async def recv_drs_pick_valuer(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """By Valuer: handle the valuer-picker tap, then show the period picker."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()

    if query.data == "drs_pickvaluer:cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    sess = _get_drs_sess(ctx)
    idx  = int(query.data.split(":")[1])
    if idx >= len(sess.valuer_choices):
        return DRS.PICK_VALUER
    valuer = sess.valuer_choices[idx]
    sess.target_key  = valuer["key"]
    sess.target_name = valuer["name"]

    await query.edit_message_text(
        f"👤 *{md_escape(valuer['name'])}*\n\nFilter history by period:",
        parse_mode="Markdown",
        reply_markup=_drs_period_keyboard(),
    )
    return DRS.PERIOD


async def recv_drs_pick_tag(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """By Tag: handle the tag-picker tap, then show the period picker."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()
    tag = query.data.split(":", 1)[1]

    if tag == "cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    sess = _get_drs_sess(ctx)
    sess.target_key  = tag
    sess.target_name = tag

    await query.edit_message_text(
        f"🏷 *{tag}*\n\nFilter history by period:",
        parse_mode="Markdown",
        reply_markup=_drs_period_keyboard(),
    )
    return DRS.PERIOD


async def recv_drs_period(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Handle the period-picker tap, then ask how often to run."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()

    if query.data == "drs_period:cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    days = int(query.data.split(":")[1])
    sess = _get_drs_sess(ctx)
    sess.period_days  = days
    sess.period_label = next((label for label, d in _DT_PERIOD_OPTIONS if d == days), "All time")

    await query.edit_message_text(
        f"✅ Period: *{sess.period_label}*\n\nHow often should this run?",
        parse_mode="Markdown",
        reply_markup=_drs_interval_keyboard(),
    )
    return DRS.INTERVAL


async def recv_drs_interval(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Handle the interval-picker tap, then ask for the recipient email."""
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()

    if query.data == "drs_interval:cancel":
        await query.edit_message_text("❌ Cancelled.")
        await query.message.reply_text("Main menu:", reply_markup=_main_menu())
        return ConversationHandler.END

    sess = _get_drs_sess(ctx)
    sess.interval_minutes = int(query.data.split(":")[1])

    await query.edit_message_text(
        f"✅ Interval: *{sess.interval_minutes} min*\n\nEnter the email address to receive this report:",
        parse_mode="Markdown",
    )
    return DRS.EMAIL


async def recv_drs_email(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Validate the entered email, save the schedule, and register its repeating job."""
    if not allowed(update): return await deny(update)
    email = update.message.text.strip()
    if not _DRS_EMAIL_RE.fullmatch(email):
        await update.message.reply_text(
            "❌ Invalid email address. Enter a valid email (e.g. `user@example.com`).",
            parse_mode="Markdown",
        )
        return DRS.EMAIL

    sess = _get_drs_sess(ctx)
    cfg = {
        "scope":            sess.scope,
        "target_key":       sess.target_key,
        "target_name":      sess.target_name,
        "period_days":      sess.period_days,
        "period_label":     sess.period_label,
        "interval_minutes": sess.interval_minutes,
        "email":            email,
    }
    schedule_id = add_dlv_report_schedule(cfg)
    ctx.job_queue.run_repeating(
        _drs_job,
        interval=sess.interval_minutes * 60,
        first=_DRS_RESTORE_FIRST_RUN_DELAY,
        name=f"drs_job:{schedule_id}",
        data=schedule_id,
    )

    await update.message.reply_text(
        f"✅ *Schedule created*\n\n{_drs_format_schedule_summary({**cfg, 'id': schedule_id})}",
        parse_mode="Markdown",
        reply_markup=_main_menu(),
    )
    return ConversationHandler.END


# ──────────────────────────────────────────────────────────
# Registration — called from bot.py's main()
# ──────────────────────────────────────────────────────────
def register(app: Application) -> None:
    """Wire the DLV Report Schedule conversation into the given
    Application, and restore one repeating job per saved schedule on
    startup."""
    drs_conv = ConversationHandler(
        entry_points=[
            CommandHandler("dlvreportschedule", cmd_drs),
            MessageHandler(filters.Regex(f"^{re.escape(BTN_DLV_REPORT_SCHEDULE)}$"), cmd_drs),
        ],
        states={
            DRS.MENU:        [CallbackQueryHandler(recv_drs_menu,        pattern=r"^drs_menu:")],
            DRS.REMOVE_PICK: [CallbackQueryHandler(recv_drs_remove,      pattern=r"^drs_remove:")],
            DRS.SCOPE:       [CallbackQueryHandler(recv_drs_scope,       pattern=r"^drs_scope:")],
            DRS.PICK_VALUER: [CallbackQueryHandler(recv_drs_pick_valuer, pattern=r"^drs_pickvaluer:")],
            DRS.PICK_TAG:    [CallbackQueryHandler(recv_drs_pick_tag,    pattern=r"^drs_picktag:")],
            DRS.PERIOD:      [CallbackQueryHandler(recv_drs_period,      pattern=r"^drs_period:")],
            DRS.INTERVAL:    [CallbackQueryHandler(recv_drs_interval,    pattern=r"^drs_interval:")],
            DRS.EMAIL:       [MessageHandler(not_cancel, recv_drs_email)],
        },
        fallbacks=[
            CommandHandler("cancel", cmd_cancel),
            MessageHandler(_CANCEL_FILTER, cmd_cancel),
            MessageHandler(filters.TEXT, fallback),
        ],
        allow_reentry=True,
        per_message=False,
    )
    app.add_handler(drs_conv)

    schedules = load_dlv_report_schedules()
    for cfg in schedules:
        schedule_id = cfg.get("id", "")
        if not schedule_id:
            continue
        app.job_queue.run_repeating(
            _drs_job,
            interval=cfg.get("interval_minutes", 60) * 60,
            first=_DRS_RESTORE_FIRST_RUN_DELAY,
            name=f"drs_job:{schedule_id}",
            data=schedule_id,
        )
    if schedules:
        logger.info(
            "DLV Report Schedule: restored %d schedule(s) (first run in %ds each)",
            len(schedules), _DRS_RESTORE_FIRST_RUN_DELAY,
        )
