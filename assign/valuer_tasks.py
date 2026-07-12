#!/usr/bin/env python3
"""
valuer_tasks.py
===============
Valuer Tasks — look up all stamp-duty tasks assigned to a specific valuer
within a date range and export the results to Excel (/valuertasks or
"👤 Valuer Tasks").

Fetches Ongoing + Completed candidates in the date range, then detail-fetches
each one (with multi-credential token rotation on 403, via token_rotator.py)
to confirm the valuer actually holds it.

Call register(app) from bot.py's main() to wire this feature in.
"""

import asyncio
import io
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed as _futures_as_completed
from dataclasses import dataclass
from datetime import datetime
from enum import Enum, auto
from typing import List

import openpyxl
import requests
from openpyxl.styles import Font, PatternFill
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

from ardhisasa_auth import AuthTokens, build_session
from common import (
    BASE_URL,
    BTN_VALUER_TASKS,
    CPARAMS_DLV,
    CRED_LABELS,
    CRED_MAP,
    _be_cred_keyboard,
    _CANCEL_FILTER,
    _date_cutoff_str,
    _main_menu,
    _NODE_LABELS,
    _within_days,
    allowed,
    cmd_cancel,
    deny,
    fallback,
    get_valid_tokens,
    logger,
    not_cancel,
)
from token_rotator import _AllTokensExhausted, _TokenRotator

_VT_TOKEN_ROTATE_DELAY = 10   # seconds to wait before retrying with a new token
_VT_MAX_RETRIES        = 5    # max 429/5xx retries before aborting a detail fetch
_VT_WORKERS            = 5


# ──────────────────────────────────────────────────────────
# States — Valuer Tasks conversation
# ──────────────────────────────────────────────────────────
class VT(Enum):
    PICK_CRED     = auto()   # select cached credential
    STAFF_NAME    = auto()   # enter valuer name to search
    SELECT_STAFF  = auto()   # pick from search results
    DAYS_BACK     = auto()   # enter / pick number of days back


@dataclass
class VTSession:
    cred_type:    str  = ""
    valuer_name:  str  = ""
    valuer_uid:   str  = ""


def _get_vt_sess(ctx: ContextTypes.DEFAULT_TYPE) -> VTSession:
    if "vt_session" not in ctx.user_data:
        ctx.user_data["vt_session"] = VTSession()
    return ctx.user_data["vt_session"]


_VT_TASK_LIST_URL = f"{BASE_URL}/valuationservice/api/v1/stamp-duty/application"
_VT_DETAIL_URL    = f"{BASE_URL}/valuationservice/api/v1/stamp-duty/application/detail-view"


def _vt_headers(tokens: AuthTokens) -> dict:
    return {
        "Authorization": f"Bearer {tokens.access_token}",
        "JWTAUTH":       f"Bearer {tokens.jwt}",
        "cparams":       CPARAMS_DLV,
    }


def _vt_fetch_task_detail(sess: requests.Session, rotator: "_TokenRotator", task_id: str) -> dict:
    """Fetch detail for one task with token rotation on 403."""
    retries = 0
    while True:
        tokens = rotator.current()
        if tokens is None:
            raise _AllTokensExhausted(f"All tokens exhausted fetching task {task_id}")
        headers = _vt_headers(tokens)
        resp = sess.get(_VT_DETAIL_URL, headers=headers, params={"request_id": task_id}, timeout=30)
        if resp.status_code == 403:
            new_tokens = rotator.rotate(tokens)
            if new_tokens is None:
                raise _AllTokensExhausted(f"All tokens exhausted (403) fetching task {task_id}")
            time.sleep(_VT_TOKEN_ROTATE_DELAY)
            continue
        if resp.status_code in (429, 502, 503, 504):
            retries += 1
            if retries >= _VT_MAX_RETRIES:
                raise RuntimeError(
                    f"_vt_fetch_task_detail: too many transient errors (HTTP {resp.status_code}) "
                    f"for task_id={task_id}"
                )
            time.sleep(2)
            continue
        resp.raise_for_status()
        return resp.json()


def _vt_fetch_all_tasks(
    sess: requests.Session,
    headers: dict,
    task_filter: str,
    cutoff: str,
) -> List[dict]:
    """Fetch all pages for one filter bucket, dropping tasks older than cutoff."""
    tasks: List[dict] = []
    page = 1
    while True:
        try:
            resp = sess.get(
                _VT_TASK_LIST_URL,
                headers=headers,
                params={
                    "filter":       task_filter,
                    "role":         "DLV",
                    "request_type": "STAMP_DUTY",
                    "search":       "",
                    "page":         page,
                },
                timeout=30,
            )
            resp.raise_for_status()
            data    = resp.json()
            results = data.get("results") or []
            for t in results:
                if _within_days(t.get("date_created", ""), cutoff):
                    tasks.append(t)
                else:
                    # Results are date-desc sorted; once we pass the cutoff stop paging
                    return tasks
            if not data.get("next"):
                break
            page += 1
        except Exception as e:
            logger.warning("VT fetch page %d filter=%s: %s", page, task_filter, e)
            break
    return tasks


def _vt_build_excel(
    valuer_name: str,
    days_back: int,
    tasks: List[dict],
) -> bytes:
    """Build a single-sheet Excel report of the valuer's tasks."""
    wb       = openpyxl.Workbook()
    ws       = wb.active
    ws.title = "Valuer Tasks"

    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="BDD7EE")
    alt_fill = PatternFill("solid", fgColor="F2F2F2")

    headers = [
        "Reference Number", "Parcel Number", "Registry", "County",
        "Consideration (KES)", "Node / Status", "Date Created",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell      = ws.cell(row=1, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes    = "A2"

    for i, t in enumerate(tasks, start=2):
        node_raw = t.get("node", "")
        node_label = _NODE_LABELS.get(node_raw, node_raw or "—")
        row = [
            t.get("reference_number", ""),
            t.get("parcel_number", ""),
            t.get("registry", ""),
            t.get("county", ""),
            t.get("consideration_amount", ""),
            node_label,
            (t.get("date_created") or "")[:10],
        ]
        ws.append(row)
        if i % 2 == 0:
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = alt_fill

    # Summary row
    ws.append([])
    ws.append([f"Valuer: {valuer_name}", "", "", "", "", "", f"Days back: {days_back}"])
    ws.append([f"Total tasks: {len(tasks)}"])

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = max(12, min(55, max_len + 2))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _vt_run(
    tokens:       AuthTokens,
    chat_id:      int,
    bot,
    loop,
    valuer_uid:   str,
    valuer_name:  str,
    days_back:    int,
) -> None:
    """Background worker: find tasks assigned to a specific valuer within days_back."""
    def _tg(text: str):
        asyncio.run_coroutine_threadsafe(
            bot.send_message(chat_id, text, parse_mode="Markdown"),
            loop,
        ).result(timeout=15)

    cutoff  = _date_cutoff_str(days_back)
    headers = {
        "Authorization": f"Bearer {tokens.access_token}",
        "JWTAUTH":       f"Bearer {tokens.jwt}",
        "cparams":       CPARAMS_DLV,
    }

    try:
        _tg(
            f"🔍 Searching tasks for *{valuer_name}* over the last *{days_back}* day(s)…\n"
            f"_(cutoff: {cutoff})_"
        )

        sess = build_session()

        # Fetch Ongoing + Completed in parallel
        with ThreadPoolExecutor(max_workers=2) as pool:
            ongoing_fut   = pool.submit(_vt_fetch_all_tasks, sess, headers, "Ongoing",   cutoff)
            completed_fut = pool.submit(_vt_fetch_all_tasks, sess, headers, "Completed", cutoff)
            ongoing_tasks   = ongoing_fut.result()
            completed_tasks = completed_fut.result()

        candidates = ongoing_tasks + completed_tasks
        _tg(
            f"📄 *{len(candidates)}* candidate task(s) in date range "
            f"({len(ongoing_tasks)} ongoing, {len(completed_tasks)} completed).\n"
            f"Checking assignments…"
        )

        if not candidates:
            _tg(f"ℹ️ No tasks found in the last *{days_back}* day(s).")
            return

        # Detail-fetch each candidate and match against valuer_uid
        token_pairs = [(ct, get_valid_tokens(ct)) for ct in CRED_MAP if get_valid_tokens(ct)]
        token_pairs.sort(key=lambda p: 0 if p[1] is tokens else 1)
        rotator    = _TokenRotator(token_pairs)
        matched    = []
        exhausted  = False

        with ThreadPoolExecutor(max_workers=_VT_WORKERS) as pool:
            fut_map = {
                pool.submit(_vt_fetch_task_detail, sess, rotator, t["id"]): t
                for t in candidates
            }
            for fut in _futures_as_completed(fut_map):
                task_summary = fut_map[fut]
                try:
                    if exhausted:
                        fut.cancel()
                        continue
                    detail = fut.result()
                    actors = detail.get("actors") or []
                    vo     = next((a for a in actors if a.get("role") == "VALUATION OFFICER"), None)
                    if vo and (vo.get("user_details") or {}).get("id", "") == valuer_uid:
                        ext = detail.get("external_process_details") or {}
                        matched.append({
                            "reference_number":    detail.get("reference_number", ""),
                            "parcel_number":       detail.get("parcel_number", ""),
                            "registry":            detail.get("registry", ""),
                            "county":              detail.get("county", ""),
                            "consideration_amount": ext.get("consideration_amount", ""),
                            "node":                detail.get("node", ""),
                            "date_created":        detail.get("date_created", ""),
                        })
                except _AllTokensExhausted:
                    exhausted = True
                    logger.warning("VT: all tokens exhausted")
                except Exception as e:
                    logger.warning("VT detail failed for %s: %s", task_summary.get("id"), e)

        if exhausted:
            _tg("⚠️ Tokens exhausted during detail fetch — results may be partial.")

        if not matched:
            _tg(
                f"ℹ️ No tasks assigned to *{valuer_name}* found in the last *{days_back}* day(s)."
            )
            return

        # Sort by date descending
        matched.sort(key=lambda t: t.get("date_created", ""), reverse=True)

        _tg(
            f"✅ *{len(matched)}* task(s) found for *{valuer_name}*.\nBuilding report…"
        )

        xlsx_bytes = _vt_build_excel(valuer_name, days_back, matched)
        filename   = (
            f"ValuerTasks_{valuer_name.replace(' ', '_')}"
            f"_{days_back}days_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

        asyncio.run_coroutine_threadsafe(
            bot.send_document(
                chat_id,
                document=io.BytesIO(xlsx_bytes),
                filename=filename,
                caption=(
                    f"👤 *{valuer_name}* — last {days_back} day(s)\n"
                    f"Tasks found: {len(matched)}"
                ),
            ),
            loop,
        ).result(timeout=60)

    except Exception as exc:
        logger.exception("VT run failed: %s", exc)
        _tg(f"❌ Valuer Tasks failed: `{exc}`")


# ── Valuer Tasks conversation handlers ────────────────────

async def cmd_valuer_tasks(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not allowed(update): return await deny(update)
    sess           = _get_vt_sess(ctx)
    sess.cred_type = ""

    kbd = _be_cred_keyboard()
    if not kbd:
        await update.message.reply_text(
            "❌ No valid cached tokens. Use *🔑 Refresh Auth* first.",
            parse_mode="Markdown",
            reply_markup=_main_menu(),
        )
        return ConversationHandler.END

    await update.message.reply_text(
        "👤 *Valuer Tasks*\n\n"
        "Look up all tasks assigned to a specific valuer within a date range.\n\n"
        "Select the account to search with:",
        parse_mode="Markdown",
        reply_markup=kbd,
    )
    return VT.PICK_CRED


async def recv_vt_cred(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()

    sess           = _get_vt_sess(ctx)
    sess.cred_type = query.data.split(":")[1]

    await query.edit_message_text(
        f"✅ Account: *{CRED_LABELS.get(sess.cred_type, sess.cred_type)}*\n\n"
        "Enter the *valuer's name* to search:",
        parse_mode="Markdown",
    )
    return VT.STAFF_NAME


async def recv_vt_name(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not allowed(update): return await deny(update)

    sess = _get_vt_sess(ctx)
    name = (update.message.text or "").strip()
    if not name:
        await update.message.reply_text("Please enter a name.")
        return VT.STAFF_NAME

    tokens = get_valid_tokens(sess.cred_type)
    if not tokens:
        await update.message.reply_text(
            "❌ Tokens expired. Use *🔑 Refresh Auth* first.",
            parse_mode="Markdown", reply_markup=_main_menu(),
        )
        return ConversationHandler.END

    await update.message.reply_text(f"🔍 Searching for *{name}*…", parse_mode="Markdown")
    try:
        http_sess = build_session()
        resp = http_sess.get(
            f"{BASE_URL}/acl/api/v1/accounts/list-user-accounts",
            headers={"Authorization": f"Bearer {tokens.access_token}", "JWTAUTH": f"Bearer {tokens.jwt}"},
            params={"account_type": "STAFF", "filter_type": "ACTIVE", "page": 1, "search": name},
            timeout=30,
        )
        resp.raise_for_status()
        results = resp.json().get("results", [])
    except Exception as e:
        await update.message.reply_text(
            f"❌ Search failed: `{e}`", parse_mode="Markdown", reply_markup=_main_menu(),
        )
        return ConversationHandler.END

    if not results:
        await update.message.reply_text(
            f"⚠️ No staff found for *{name}*. Try a different name.",
            parse_mode="Markdown",
        )
        return VT.STAFF_NAME

    ctx.user_data["vt_search_results"] = results
    rows = []
    for i, v in enumerate(results):
        sd         = v.get("staff_details", {})
        full_name  = " ".join(filter(None, [sd.get("firstname"), sd.get("middlename"), sd.get("lastname")]))
        rows.append([InlineKeyboardButton(full_name or f"Staff {i+1}", callback_data=f"vt_staff:{i}")])

    await update.message.reply_text(
        f"Found *{len(results)}* result(s). Select a valuer:",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(rows),
    )
    return VT.SELECT_STAFF


async def recv_vt_select(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()

    idx     = int(query.data.split(":")[1])
    results = ctx.user_data.get("vt_search_results", [])
    if idx >= len(results):
        await query.edit_message_text("❌ Invalid selection.")
        await ctx.bot.send_message(query.message.chat_id, "Main menu.", reply_markup=_main_menu())
        return ConversationHandler.END

    v   = results[idx]
    sd  = v.get("staff_details", {})
    uid = sd.get("user_id") or v.get("id", "")
    full_name = " ".join(filter(None, [sd.get("firstname"), sd.get("middlename"), sd.get("lastname")]))

    sess              = _get_vt_sess(ctx)
    sess.valuer_name  = full_name
    sess.valuer_uid   = uid

    await query.edit_message_text(
        f"✅ Valuer: *{full_name}*\n\nHow many days back do you want to check?",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup([
            [
                InlineKeyboardButton("7 days",  callback_data="vt_days:7"),
                InlineKeyboardButton("14 days", callback_data="vt_days:14"),
                InlineKeyboardButton("30 days", callback_data="vt_days:30"),
            ],
            [
                InlineKeyboardButton("60 days",  callback_data="vt_days:60"),
                InlineKeyboardButton("90 days",  callback_data="vt_days:90"),
                InlineKeyboardButton("Custom…",  callback_data="vt_days:custom"),
            ],
        ]),
    )
    return VT.DAYS_BACK


async def recv_vt_days_callback(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not allowed(update): return await deny(update)
    query = update.callback_query
    await query.answer()

    value = query.data.split(":")[1]
    if value == "custom":
        await query.edit_message_text(
            "Enter the number of days back (e.g. `45`):",
            parse_mode="Markdown",
        )
        return VT.DAYS_BACK

    return await _vt_start_run(query.message.chat_id, ctx, int(value), edit_msg=query)


async def recv_vt_days_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not allowed(update): return await deny(update)
    raw = (update.message.text or "").strip()
    try:
        days = int(raw)
        if days < 1:
            raise ValueError
    except ValueError:
        await update.message.reply_text("Please enter a positive whole number (e.g. `30`).", parse_mode="Markdown")
        return VT.DAYS_BACK

    return await _vt_start_run(update.message.chat_id, ctx, days, reply_msg=update.message)


async def _vt_start_run(
    chat_id: int,
    ctx: ContextTypes.DEFAULT_TYPE,
    days_back: int,
    edit_msg=None,
    reply_msg=None,
):
    sess   = _get_vt_sess(ctx)
    tokens = get_valid_tokens(sess.cred_type)
    if not tokens:
        text = f"❌ Tokens expired. Use *🔑 Refresh Auth* first."
        if edit_msg:
            await edit_msg.edit_text(text, parse_mode="Markdown")
        else:
            await reply_msg.reply_text(text, parse_mode="Markdown")
        await ctx.bot.send_message(chat_id, "Main menu.", reply_markup=_main_menu())
        return ConversationHandler.END

    confirm_text = (
        f"⏳ Fetching tasks for *{sess.valuer_name}* — last *{days_back}* day(s)…\n"
        "You will be notified when done."
    )
    if edit_msg:
        await edit_msg.edit_text(confirm_text, parse_mode="Markdown")
    else:
        await reply_msg.reply_text(confirm_text, parse_mode="Markdown")
    await ctx.bot.send_message(chat_id, "Returning to menu.", reply_markup=_main_menu())

    loop = asyncio.get_event_loop()
    asyncio.ensure_future(
        asyncio.to_thread(
            _vt_run, tokens, chat_id, ctx.bot, loop,
            sess.valuer_uid, sess.valuer_name, days_back,
        )
    )
    return ConversationHandler.END


# ──────────────────────────────────────────────────────────
# Registration
# ──────────────────────────────────────────────────────────

def register(app: Application) -> None:
    """Wire the Valuer Tasks conversation into the given Application."""
    vt_conv = ConversationHandler(
        entry_points=[
            CommandHandler("valuertasks", cmd_valuer_tasks),
            MessageHandler(filters.Regex(f"^{re.escape(BTN_VALUER_TASKS)}$"), cmd_valuer_tasks),
        ],
        states={
            VT.PICK_CRED:    [CallbackQueryHandler(recv_vt_cred,          pattern=r"^be_cred:")],
            VT.STAFF_NAME:   [MessageHandler(not_cancel, recv_vt_name)],
            VT.SELECT_STAFF: [CallbackQueryHandler(recv_vt_select,        pattern=r"^vt_staff:")],
            VT.DAYS_BACK:    [
                CallbackQueryHandler(recv_vt_days_callback, pattern=r"^vt_days:"),
                MessageHandler(not_cancel, recv_vt_days_text),
            ],
        },
        fallbacks=[
            CommandHandler("cancel", cmd_cancel),
            MessageHandler(_CANCEL_FILTER, cmd_cancel),
            MessageHandler(filters.TEXT, fallback),
        ],
        allow_reentry=True,
        per_message=False,
    )
    app.add_handler(vt_conv)
